"""Arabic -> English translation of submission workbooks, driven by a user-supplied dictionary
(column A = Arabic, column B = English, header in row 1).

This is the supplied translation script, turned into importable functions: the same
dictionary loading, text normalisation, whole-cell / word-by-word replacement, sheet-name
translation, formula sheet-reference rewriting and "missing words" report - minus the
tkinter file pickers (a server has no display) and with xlrd / pyxlsb imported lazily so
they are only needed when an .xls / .xlsb file is actually uploaded.

Python 3.9 compatible.
"""
import os
import re
import warnings

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

warnings.filterwarnings("ignore", message="Cannot parse header or footer so it will be ignored")

ARABIC_PATTERN = re.compile(r"[؀-ۿݐ-ݿࢠ-ࣿ]")
ARABIC_WORD_PATTERN = re.compile(r"[؀-ۿݐ-ݿࢠ-ࣿ]+")

SUPPORTED_EXTENSIONS = (".xlsx", ".xlsm", ".xls", ".xlsb")


def contains_arabic(text):
    return isinstance(text, str) and bool(ARABIC_PATTERN.search(text))


# ----------------------------------------------------------------------------- normalisation

def normalize_text(text):
    if text is None:
        return ""
    text = str(text).strip()
    text = text.replace("*", "")           # bullets
    text = text.replace("•", "")
    text = text.replace("ـ", "")           # tatweel
    text = text.replace("أ", "ا").replace("إ", "ا").replace("آ", "ا").replace("ٱ", "ا")
    text = text.replace("ى", "ي")
    text = text.replace("ة", "ه")
    text = re.sub(r"[ً-ٰٟۖ-ۭ]", "", text)   # diacritics
    text = re.sub(r"\s+", " ", text)
    return text.strip()


# ----------------------------------------------------------------------------- dictionary

def load_dictionary(dictionary_file):
    """Column A = Arabic source, column B = English target; row 1 is a header. Keys are stored in
    normalised form so spelling variants of the same Arabic word hit the same entry."""
    wb = load_workbook(dictionary_file, data_only=True)
    ws = wb.active
    dictionary = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        source = row[0]
        target = row[1] if len(row) > 1 else None
        if source and target:
            dictionary[normalize_text(source)] = str(target).strip()
    return dictionary


# ----------------------------------------------------------------------------- helpers

def apply_ltr(cell):
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=False, readingOrder=1, indent=0)


def normalize_sheet_title(title):
    """Excel: max 31 chars, none of  : \\ / ? * [ ]"""
    cleaned = re.sub(r"[:\\/?*\[\]]", " ", str(title))
    cleaned = re.sub(r"\s+", " ", cleaned).strip()
    return (cleaned or "Sheet")[:31]


def make_unique_sheet_title(base_title, used_titles):
    candidate = base_title
    if candidate not in used_titles:
        return candidate
    index = 2
    while True:
        suffix = " (%d)" % index
        trimmed = base_title[:31 - len(suffix)].rstrip()
        candidate = "%s%s" % (trimmed, suffix)
        if candidate not in used_titles:
            return candidate
        index += 1


def _quote_sheet_name(name):
    return "'" + str(name).replace("'", "''") + "'"


def update_formula_sheet_references(formula, title_map):
    """Rewrite sheet references inside a formula after tabs were renamed (quoted and unquoted)."""
    if not isinstance(formula, str) or not formula.startswith("="):
        return formula
    updated = formula
    for old_title, new_title in title_map.items():
        if old_title == new_title:
            continue
        updated = updated.replace("%s!" % _quote_sheet_name(old_title), "%s!" % _quote_sheet_name(new_title))
        if re.match(r"^[A-Za-z0-9_.]+$", old_title):
            pattern = re.compile(r"(?<![A-Za-z0-9_\.])" + re.escape(old_title) + r"!")
            updated = pattern.sub("%s!" % new_title, updated)
    return updated


def _merged_lookup(ws):
    """coordinate -> (master_coordinate) for every merged cell; computed once per sheet, because
    asking every cell to scan every merged range is quadratic on a large sheet."""
    lookup = {}
    for rng in ws.merged_cells.ranges:
        master = ws.cell(rng.min_row, rng.min_col).coordinate
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                lookup[ws.cell(r, c).coordinate] = master
    return lookup


# ----------------------------------------------------------------------------- translation

def collapse_letter_hyphens(text):
    """'t-r-a-n-s-l-a-t-i-o-n' -> 'translation'"""
    if not isinstance(text, str) or "-" not in text:
        return text
    letter_chain = re.compile(r"\b(?:[A-Za-z]\s*-\s*){2,}[A-Za-z]\b")
    return letter_chain.sub(lambda m: re.sub(r"\s*-\s*", "", m.group(0)), text)


def translate_text(text, dictionary, missing_terms):
    original = str(text)
    normalized_original = normalize_text(original)

    # the whole cell is a dictionary entry
    if normalized_original in dictionary:
        return collapse_letter_hyphens(dictionary[normalized_original])

    translated_parts = []
    last_end = 0
    translated_any = False
    for match in ARABIC_WORD_PATTERN.finditer(original):
        start, end = match.span()
        arabic_word = match.group(0)
        translated_parts.append(original[last_end:start])
        normalized_word = normalize_text(arabic_word)
        if normalized_word in dictionary:
            translated_parts.append(dictionary[normalized_word])
            translated_any = True
        else:
            translated_parts.append(arabic_word)
            missing_terms.add(arabic_word)
        last_end = end
    translated_parts.append(original[last_end:])

    if not translated_any:
        return original
    return collapse_letter_hyphens("".join(translated_parts))


def write_missing_terms(missing_terms, out_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Missing Arabic Words"
    ws["A1"] = "Arabic Text Not Found"
    for i, word in enumerate(sorted(missing_terms), start=2):
        ws.cell(row=i, column=1, value=word)
    ws.column_dimensions["A"].width = 60
    wb.save(out_path)
    return out_path


# ----------------------------------------------------------------------------- .xls / .xlsb -> openpyxl

_XLS_BORDER_STYLES = {0: None, 1: "thin", 2: "medium", 3: "dashed", 4: "dotted", 5: "thick", 6: "double", 7: "hair",
                      8: "mediumDashed", 9: "dashDot", 10: "mediumDashDot", 11: "dashDotDot",
                      12: "mediumDashDotDot", 13: "slantDashDot"}
_XLS_H_ALIGN = {0: "general", 1: "left", 2: "center", 3: "right", 4: "fill", 5: "justify", 6: "centerContinuous", 7: "distributed"}
_XLS_V_ALIGN = {0: "top", 1: "center", 2: "bottom", 3: "justify", 4: "distributed"}


def _xls_rgb(colour_map, index):
    rgb = colour_map.get(index)
    if rgb:
        return "{:02X}{:02X}{:02X}".format(*rgb)
    return None


def _xls_side(line_type, colour_index, colour_map):
    style = _XLS_BORDER_STYLES.get(line_type)
    if not style:
        return Side()
    return Side(style=style, color=_xls_rgb(colour_map, colour_index) or "000000")


def xls_to_openpyxl(source_file):
    """Old .xls via xlrd, keeping fonts, fills, borders, alignment, merges, widths and heights."""
    try:
        import xlrd
    except ImportError:
        raise RuntimeError("Reading .xls files needs the 'xlrd' package in the Python environment "
                           "(pip install xlrd). Or save the file as .xlsx and upload that.")
    xls_wb = xlrd.open_workbook(source_file, formatting_info=True)
    wb = Workbook()
    wb.remove(wb.active)
    for sheet_name in xls_wb.sheet_names():
        xls_ws = xls_wb.sheet_by_name(sheet_name)
        ws = wb.create_sheet(title=sheet_name[:31])
        for rlo, rhi, clo, chi in xls_ws.merged_cells:
            ws.merge_cells(start_row=rlo + 1, start_column=clo + 1, end_row=rhi, end_column=chi)
        for col_idx in range(xls_ws.ncols):
            info = xls_ws.colinfo_map.get(col_idx)
            if info and info.width:
                ws.column_dimensions[get_column_letter(col_idx + 1)].width = info.width / 256
        for row_idx in range(xls_ws.nrows):
            info = xls_ws.rowinfo_map.get(row_idx)
            if info and info.height:
                ws.row_dimensions[row_idx + 1].height = info.height / 20
        merged_children = set()
        for merge in ws.merged_cells.ranges:
            for r in range(merge.min_row, merge.max_row + 1):
                for c in range(merge.min_col, merge.max_col + 1):
                    if not (r == merge.min_row and c == merge.min_col):
                        merged_children.add((r, c))
        for row_idx in range(xls_ws.nrows):
            for col_idx in range(xls_ws.ncols):
                r, c = row_idx + 1, col_idx + 1
                if (r, c) in merged_children:
                    continue
                cell = xls_ws.cell(row_idx, col_idx)
                ws_cell = ws.cell(row=r, column=c, value=cell.value)
                xf_idx = cell.xf_index
                if xf_idx is None:
                    continue
                xf = xls_wb.xf_list[xf_idx]
                try:
                    f = xls_wb.font_list[xf.font_index]
                    fc = _xls_rgb(xls_wb.colour_map, f.colour_index) or "000000"
                    ws_cell.font = Font(name=f.name, size=f.height / 20, bold=bool(f.bold), italic=bool(f.italic),
                                        underline="single" if f.underline_type else None, color=fc)
                except Exception:
                    pass
                try:
                    bg = xf.background
                    if bg.pattern_type:
                        fg_hex = _xls_rgb(xls_wb.colour_map, bg.pattern_colour_index)
                        if fg_hex:
                            ws_cell.fill = PatternFill(fill_type="solid", fgColor=fg_hex)
                except Exception:
                    pass
                try:
                    b = xf.border
                    ws_cell.border = Border(
                        left=_xls_side(b.left_line_type, b.left_colour_index, xls_wb.colour_map),
                        right=_xls_side(b.right_line_type, b.right_colour_index, xls_wb.colour_map),
                        top=_xls_side(b.top_line_type, b.top_colour_index, xls_wb.colour_map),
                        bottom=_xls_side(b.bottom_line_type, b.bottom_colour_index, xls_wb.colour_map))
                except Exception:
                    pass
                try:
                    a = xf.alignment
                    ws_cell.alignment = Alignment(horizontal=_XLS_H_ALIGN.get(a.hor_align, "general"),
                                                  vertical=_XLS_V_ALIGN.get(a.vert_align, "bottom"),
                                                  wrap_text=bool(a.text_wrapped), text_rotation=a.rotation)
                except Exception:
                    pass
                try:
                    fmt = xls_wb.format_map.get(xf.format_key)
                    if fmt and fmt.format_str:
                        ws_cell.number_format = fmt.format_str
                except Exception:
                    pass
    return wb


def xlsb_to_openpyxl(source_file):
    """Binary .xlsb via pyxlsb: values and structure only (pyxlsb exposes little styling)."""
    try:
        import pyxlsb
    except ImportError:
        raise RuntimeError("Reading .xlsb files needs the 'pyxlsb' package in the Python environment "
                           "(pip install pyxlsb). Or save the file as .xlsx and upload that.")
    wb = Workbook()
    wb.remove(wb.active)
    with pyxlsb.open_workbook(source_file) as xlsb_wb:
        for sheet_name in xlsb_wb.sheets:
            ws = wb.create_sheet(title=sheet_name[:31])
            with xlsb_wb.get_sheet(sheet_name) as xlsb_ws:
                for row in xlsb_ws.rows():
                    for cell in row:
                        if cell.v is not None:
                            ws.cell(row=cell.r, column=cell.c, value=cell.v)
    return wb


def to_xlsx(source_file, out_dir):
    """.xls / .xlsb cannot be read by the rest of the app, so convert them to .xlsx first.
    .xlsx / .xlsm are returned unchanged."""
    ext = os.path.splitext(source_file)[1].lower()
    if ext not in (".xls", ".xlsb"):
        return source_file
    wb = xls_to_openpyxl(source_file) if ext == ".xls" else xlsb_to_openpyxl(source_file)
    out = os.path.join(out_dir, os.path.splitext(os.path.basename(source_file))[0] + "_converted.xlsx")
    wb.save(out)
    return out


# ----------------------------------------------------------------------------- detection + workflow

def count_arabic(path, max_cells=400000):
    """How much Arabic does a workbook hold? -> {'cells', 'sheets': [sheet names with Arabic text],
    'sheet_names': [Arabic-named sheets], 'sample': [a few strings]}. Used to decide whether to offer
    translation."""
    wb = load_workbook(path, read_only=True, data_only=True)
    cells, sheets, names, sample, seen = 0, [], [], [], 0
    for ws in wb.worksheets:
        had = False
        if contains_arabic(ws.title):
            names.append(ws.title)
        for row in ws.iter_rows(values_only=True):
            for v in row:
                seen += 1
                if isinstance(v, str) and contains_arabic(v):
                    cells += 1
                    had = True
                    if len(sample) < 5:
                        sample.append(v.strip()[:60])
            if seen >= max_cells:
                break
        if had or contains_arabic(ws.title):
            sheets.append(ws.title)
        if seen >= max_cells:
            break
    wb.close()
    return {"cells": cells, "sheets": sheets, "sheet_names": names, "sample": sample}


def translate_workbook(source_file, dictionary, output_file, keep_formulas=False):
    """Translate every Arabic text cell and sheet name of `source_file` with `dictionary` (a dict from
    load_dictionary, or a path to the dictionary workbook) and save the result to `output_file`.

    keep_formulas=False turns each formula into its saved result where one exists (openpyxl cannot
    recalculate, so a rewritten workbook would otherwise lose every total the rest of the app reads).
    Returns a stats dict including the set of Arabic terms the dictionary did not contain."""
    if not isinstance(dictionary, dict):
        dictionary = load_dictionary(dictionary)
    flattened = 0
    ext = os.path.splitext(source_file)[1].lower()
    if ext == ".xls":
        wb = xls_to_openpyxl(source_file)
    elif ext == ".xlsb":
        wb = xlsb_to_openpyxl(source_file)
    else:
        wb = load_workbook(source_file, keep_vba=(ext == ".xlsm"))
        if not keep_formulas:
            # openpyxl cannot recalculate, so a saved copy would lose every total. Replace a
            # formula with its saved result when there is one; a formula Excel never calculated
            # (no saved result) stays a formula rather than being erased.
            cached = load_workbook(source_file, data_only=True, keep_vba=(ext == ".xlsm"))
            for ws in wb.worksheets:
                wsc = cached[ws.title]
                for row in ws.iter_rows():
                    for cell in row:
                        if cell.data_type == "f":
                            val = wsc[cell.coordinate].value
                            if val is not None:
                                cell.value = val
                                flattened += 1

    missing_terms = set()
    translated_count = 0
    translated_sheet_names = 0
    used_titles = set()
    title_map = {}

    for ws in wb.worksheets:
        original_title = ws.title
        translated_title = translate_text(original_title, dictionary, missing_terms)
        unique_title = make_unique_sheet_title(normalize_sheet_title(translated_title), used_titles)
        title_map[original_title] = unique_title
        if unique_title != original_title:
            ws.title = unique_title
            translated_sheet_names += 1
        used_titles.add(ws.title)

        ws.sheet_view.rightToLeft = False        # keep the worksheet left-to-right
        merged = _merged_lookup(ws)

        for row in ws.iter_rows():
            for cell in row:
                master = merged.get(cell.coordinate)
                if master is not None and master != cell.coordinate:
                    continue                      # merged child
                value = cell.value
                if value is None or not isinstance(value, str):
                    continue
                if value.startswith("="):
                    continue
                if not contains_arabic(value):
                    continue
                translated_value = translate_text(value, dictionary, missing_terms)
                if translated_value != value:
                    cell.value = translated_value
                    apply_ltr(cell)
                    translated_count += 1

    formula_updates = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    new_formula = update_formula_sheet_references(cell.value, title_map)
                    if new_formula != cell.value:
                        cell.value = new_formula
                        formula_updates += 1

    wb.save(output_file)
    return {"translated_sheets": translated_sheet_names, "translated_cells": translated_count,
            "updated_formulas": formula_updates, "flattened_formulas": flattened, "missing_terms": missing_terms, "output_file": output_file,
            "sheet_titles": title_map}
