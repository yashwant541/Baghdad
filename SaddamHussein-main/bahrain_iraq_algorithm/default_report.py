"""Default-mapping run and its outputs.

* expand_templates   - a template of type FC_COLUMNS ("by currency") becomes one rule per currency column that the
                       target sheet really has, plus one rule that clubs every other foreign currency against the
                       'Other Foreign Currencies' column.
* run_default_mapping- applies every rule to the Trial Balance + extracted submissions, with TB amount, submission
                       amount, variance and status for each.
* write_default_workbook - per submission file: the original sheets with matched cells highlighted, the TB pivot, and a
                       report that states each rule and its variance (same idea as the depth-search output).
* build_pdf          - everything Mapping Studio shows, as a downloadable PDF.

Python 3.9 compatible.
"""
import datetime
import io
import json
import os
import re
import zipfile
from collections import OrderedDict, defaultdict

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink

from .depth_search import PALETTE, _col_title, _fill, _sheet_ref, _unique_name, build_layout
from .engine import FRX_CODE, LOCAL_CCY, text_similarity
from .pdf_simple import SimplePDF, clean

STATUS_TEXT = {"MATCH": "Match", "MATCH_WITHIN_TOLERANCE": "Match (within tolerance)",
               "REVIEW_REQUIRED": "Variance - review", "NOT_FOUND": "Not found"}
STATUS_HEX = {"MATCH": "C6EFCE", "MATCH_WITHIN_TOLERANCE": "E2EFDA", "REVIEW_REQUIRED": "F8CBAD", "NOT_FOUND": "FFEB9C"}
STATUS_RGB = {"MATCH": (0.78, 0.94, 0.81), "MATCH_WITHIN_TOLERANCE": (0.89, 0.94, 0.85),
              "REVIEW_REQUIRED": (0.97, 0.80, 0.68), "NOT_FOUND": (1.0, 0.92, 0.61)}


# ------------------------------------------------------------------------------------------ rule wording

def _friendly(patterns):
    out = []
    for p in patterns or []:
        pl = p.lower()
        if "depr" in pl or "dpr" in pl:
            w = "depreciation (Depr / Dpr)"
        elif "software" in pl or "s/w" in pl:
            w = "Capitalised Software"
        else:
            w = re.sub(r"\\s[*+]?", " ", p).replace("\\", "")
        if w not in out:
            out.append(w)
    return out


def _ccy_text(c, excl=None):
    if c == "IQD" or c == LOCAL_CCY:
        return "IQD (local currency)"
    if c == FRX_CODE:
        t = "all non-IQD currencies (forex)"
        if excl:
            t += ", except " + ", ".join(excl) + " (those have their own column)"
        return t
    if c == "TOTAL":
        return "all currencies"
    return c


def describe_group_item(it):
    bsm = str(it.get("bs_mapping") or "")
    name = ("every BS group starting with '%s' (all Assets)" % bsm[:-1]) if bsm.endswith("*") else "BS group %s" % bsm
    bits = [name, _ccy_text(it.get("currency"), it.get("exclude_currencies"))]
    if it.get("sign_filter") == "positive":
        bits.append("positive account balances only")
    elif it.get("sign_filter") == "negative":
        bits.append("negative account balances only")
    if it.get("include_desc"):
        bits.append("only accounts described as " + ", ".join(_friendly(it["include_desc"])))
    if it.get("exclude_desc"):
        bits.append("excluding " + ", ".join(_friendly(it["exclude_desc"])))
    return " | ".join(bits)


def _sub_col_text(ccy):
    if ccy == "IQD":
        return "the IQD column of that row that holds the amount (Residents or Non-Residents)"
    if ccy == FRX_CODE:
        return "the foreign-currency column of that row that holds the amount (Residents or Non-Residents)"
    if ccy == "TOTAL":
        return "Total column of that row"
    return "%s column of that row" % ccy


def describe_rule(t):
    if t.get("rule_type") == "FC_COLUMNS":
        sub = (t.get("submission_side") or [{}])[0]
        return ("BS group%s %s: each currency that has its own column on '%s' is matched to that column; every other "
                "foreign currency is clubbed together and matched to 'Other Foreign Currencies'"
                % ("s" if len(t.get("tb_groups", [])) > 1 else "", " + ".join(t.get("tb_groups", [])), sub.get("sheet")))
    groups = [describe_group_item(x) for x in t.get("tb_side", []) if x.get("level") == "group"]
    accts = [x for x in t.get("tb_side", []) if x.get("level") != "group"]
    parts = list(groups)
    if accts:
        uniq = OrderedDict()
        for x in accts:
            uniq[(x.get("bs_mapping"), x.get("account"), x.get("account_desc"), x.get("currency"))] = 1
        shown = "; ".join("%s %s %s" % (k[0], k[1], k[2]) for k in list(uniq)[:3])
        parts.append("fixed accounts (%d): %s%s" % (len(uniq), shown, " ..." if len(uniq) > 3 else ""))
    tb_text = "  +  ".join(parts) if parts else "(no Trial Balance side)"
    sub = (t.get("submission_side") or [{}])[0]
    sg = sub.get("sign", 1)
    return "%s   ->   %s > %s  [%s%s]" % (tb_text, sub.get("sheet"), sub.get("match_text"), _sub_col_text(sub.get("currency")),
                                            ", sign reversed" if sg == -1 else "")


# ------------------------------------------------------------------------------------------ expansion

def _fc_expand(t, sub):
    """One FC_COLUMNS template -> a rule per direct currency column of the sheet + an 'Other Foreign Currencies' rule."""
    out = []
    s0 = (t.get("submission_side") or [{}])[0]
    sheet = s0.get("sheet")
    on_sheet = sub[sub.sheet == sheet] if (sub is not None and len(sub)) else sub
    codes = sorted(set(on_sheet.currency)) if (on_sheet is not None and len(on_sheet)) else []
    direct = [c for c in codes if c not in ("TOTAL", "IQD", FRX_CODE, "UNSPECIFIED") and not str(c).startswith("TOTAL_")]
    has_other = FRX_CODE in codes
    if not direct and not has_other:
        out.append({"label": t["label"], "currency": "TOTAL", "rule_type": "GROUP_RULE", "tb_side": [],
                    "submission_side": copy_sub(s0, "TOTAL"), "_fc_missing": True, "derived": t.get("derived")})
        return out
    for c in direct:
        out.append({"label": "%s - %s" % (t["label"], c), "currency": c, "rule_type": "GROUP_RULE", "_expanded": True,
                    "derived": t.get("derived"), "tb_side": [grp(g, c) for g in t.get("tb_groups", [])],
                    "submission_side": copy_sub(s0, c)})
    if has_other:
        out.append({"label": "%s - Other Foreign Currencies" % t["label"], "currency": FRX_CODE, "rule_type": "GROUP_RULE",
                    "_expanded": True, "derived": t.get("derived"),
                    "tb_side": [grp(g, FRX_CODE, exclude_currencies=list(direct)) for g in t.get("tb_groups", [])],
                    "submission_side": copy_sub(s0, FRX_CODE)})
    return out


_KINDS = {"ASSET_GROUP": "Assets", "LIABILITY_GROUP": "Liabilities"}


def _sheet_names(kind, sub, override=None):
    """Names of the main / foreign-currency / maturity sheet of a statement ('Assets' or 'Liabilities'), read from the
    extracted sheets when they can be told apart, else the CBI default names."""
    names = {"main": kind, "foreign": kind + " by Foreign Currency", "maturity": kind + " by Maturity"}
    have = sorted(set(sub.sheet)) if (sub is not None and len(sub)) else []
    k = kind.lower()
    for sh in have:
        low = str(sh).lower().strip()
        if not low.startswith(k):
            continue
        if "matur" in low or "period" in low or "tenor" in low:      # 'Liabilities According to Period'
            names["maturity"] = sh
        elif "foreign" in low or "curren" in low:                    # 'Liabilities According to Curren(cy)' - Excel cuts names at 31 characters
            names["foreign"] = sh
        elif low == k:
            names["main"] = sh
    names.update(override or {})
    return names, set(have)


def _asset_group(t, sub=None, notes=None):
    """ASSET_GROUP / LIABILITY_GROUP template (the logic of 'Iraq Changes.docx'; the same for both statements): for the
    row of one BS-mapping group (or groups)
    - main sheet, IQD column 3 or 4 = TB group, IQD only (the subtotal is in ONE of them, not their sum)
    - main sheet, foreign column 5 or 6 = TB group, every non-IQD currency
    - main sheet, total column 2   = TB group, all currencies
    - '<kind> by Foreign Currency'  = TB group currency by currency; the rest against 'Other Foreign Currencies'
    - '<kind> by Maturity', total column = TB group, all currencies
    Fields: label, tb_groups, row (text of the row; row_foreign / row_maturity override it per sheet), include (subset of
    IQD, FRX, TOTAL, FC, MATURITY), sign (-1 for liabilities, whose TB balances are negative), sheets (override names).
    A foreign-currency / maturity sheet that was not extracted is skipped and mentioned in notes."""
    kind = _KINDS.get(t.get("rule_type"), "Assets")
    groups = t.get("tb_groups") or []
    names, have = _sheet_names(kind, sub, t.get("sheets"))
    inc = set(t.get("include") or ["IQD", "FRX", "TOTAL", "FC", "MATURITY"])
    sign = t.get("sign", -1 if kind == "Liabilities" else 1)
    text = t.get("row") or t.get("label")

    def side(sheet, txt, ccy):
        return [{"match_text": txt, "sheet": sheet, "currency": ccy, "is_total": False, "sign": sign}]

    out = []
    for ccy, what in (("IQD", "IQD (%s column 3 or 4)" % kind), ("FRX", "foreign currencies (%s column 5 or 6)" % kind),
                      ("TOTAL", "total (%s column 2)" % kind)):
        if ccy in inc:
            out.append({"label": "%s - %s" % (t["label"], what), "currency": ccy, "rule_type": "GROUP_RULE", "_expanded": True, "derived": True,
                        "tb_side": [grp(g, ccy) for g in groups], "submission_side": side(names["main"], text, ccy)})
    for key, label, kindname in (("FC", "by currency", "foreign"), ("MATURITY", "total (%s by Maturity)" % kind, "maturity")):
        if key not in inc:
            continue
        if have and names[kindname] not in have:
            if notes is not None and ("sheet", names[kindname]) not in notes:
                notes.append(("sheet", names[kindname]))
            continue
        if key == "FC":
            out.append({"label": "%s - %s" % (t["label"], label), "currency": "FC", "rule_type": "FC_COLUMNS", "derived": True,
                        "tb_groups": list(groups), "submission_side": side(names["foreign"], t.get("row_foreign") or text, "TOTAL")})
        else:
            out.append({"label": "%s - %s" % (t["label"], label), "currency": "TOTAL", "rule_type": "GROUP_RULE", "_expanded": True,
                        "derived": True, "tb_side": [grp(g, "TOTAL") for g in groups],
                        "submission_side": side(names["maturity"], t.get("row_maturity") or text, "TOTAL")})
    return out


def _rule_key(t):
    tb = []
    for x in t.get("tb_side") or []:
        if x.get("level") != "group":
            return None
        tb.append((str(x.get("bs_mapping")).lower(), x.get("currency"), x.get("sign_filter"), tuple(x.get("include_desc") or ()),
                   tuple(x.get("exclude_desc") or ()), tuple(sorted(x.get("exclude_currencies") or ())), x.get("sign", 1)))
    if not tb:
        return None
    subs = tuple((x.get("sheet"), re.sub(r"\W+", "", str(x.get("match_text")).lower())[:30], x.get("currency"), x.get("sign", 1))
                 for x in t.get("submission_side") or [])
    return (tuple(sorted(tb)), subs)


def expand_templates(templates, tb, sub, notes=None):
    """Concrete templates. ASSET_GROUP and FC_COLUMNS templates expand from the target sheets' own currency columns;
    a derived (ASSET_GROUP) rule that an explicit rule already covers is dropped, so nothing is counted twice."""
    staged = []
    for t in templates:
        if t.get("rule_type") in _KINDS:
            staged.extend(_asset_group(t, sub, notes))
        else:
            staged.append(t)
    out = []
    for t in staged:
        if t.get("rule_type") == "FC_COLUMNS":
            out.extend(_fc_expand(t, sub))
        else:
            out.append(t)
    explicit = set(k for k in (_rule_key(t) for t in out if not t.get("derived")) if k)
    return [t for t in out if not (t.get("derived") and _rule_key(t) in explicit)]


def grp(bsm, ccy, **kw):
    d = {"level": "group", "bs_mapping": bsm, "currency": ccy, "sign": 1}
    d.update(kw)
    return d


def copy_sub(s0, ccy):
    d = dict(s0)
    d["currency"] = ccy
    return [d]


# ------------------------------------------------------------------------------------------ off-balance checks
# The same basic rules as Assets / Liabilities, but the values come from the Outstanding Report pivots:
#   main sheet ('Off-Balance Sheet Accounts')      IQD column 3 or 4 / foreign column 5 or 6 / total column 2  <- Pivot 1 (rows LC/GTEE DESC, columns BILL_CCY)
#   'Off-Balance Sheet Accounts by F(oreign ...)'  currency by currency, the rest against 'Other'             <- Pivot 1
#   'Off-Balance Sheet Accounts by M(aturity)'     one check per maturity bucket column, plus the total       <- Pivot 2 (rows CATEGORY > LC/GTEE > BILL_CCY, columns Bucket)

OB_SOURCE = "Outstanding Report"
_NUMWORDS = {"one": "1", "two": "2", "three": "3", "four": "4", "five": "5", "six": "6", "seven": "7", "eight": "8", "nine": "9",
             "ten": "10", "eleven": "11", "twelve": "12"}


def _ob_sheet_names(sub):
    have = sorted(set(sub.sheet)) if (sub is not None and len(sub)) else []
    main, foreign, maturity = [], None, None
    for sh in have:
        low = re.sub(r"\s+", " ", re.sub(r"[^a-z0-9 ]", " ", str(sh).lower())).strip()
        if not (low.startswith("off") and "balance" in low):
            continue
        if "matur" in low or "period" in low or "tenor" in low or low.endswith(" by m"):
            maturity = sh                                              # 'Off-Balance Sheet Accounts by M' (name cut at 31 characters)
        elif "foreign" in low or "curren" in low or low.endswith(" by f"):
            foreign = sh
        else:
            main.append(sh)
    return {"main": min(main, key=len) if main else None, "foreign": foreign, "maturity": maturity}


def _norm_item(s):
    """'GOODS/SERVICE G'TEE' and 'goods service g tee' are the same item."""
    return re.sub(r"[^a-z0-9]", "", str(s).lower())


def load_ob_row_groups():
    """Sheet rows that gather several LC/GTEE items, e.g. 'For Other Purposes' = ADVANCE PAYMENT BOND + LTR OF INTENT + ...
    (default_mapping.json -> "ob_row_groups": [{"row": ..., "items": [...]}])."""
    try:
        path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "default_mapping.json")
        with open(path, encoding="utf-8") as fh:
            data = json.load(fh)
        return [g for g in (data.get("ob_row_groups") or []) if g.get("row") and g.get("items")]
    except (OSError, ValueError):
        return []


def _sheet_columns(sub, sheet):
    """[(column letter, header text)] of the extracted cells of a sheet, left to right."""
    out = {}
    d = sub[sub.sheet == sheet]
    for r in d.itertuples(index=False):
        m = re.match(r"^([A-Z]+)\d+$", str(r.source_cell))
        if m and m.group(1) not in out:
            out[m.group(1)] = str(getattr(r, "column_context", "") or "").split(" | ")[0].strip()      # drop the sample value the context carries
    return [(L, out[L]) for L in sorted(out, key=lambda L: (len(L), L))]


def _upper_bound_months(label):
    """Upper bound, in months, of a maturity label ('1M', '0-30 Days', 'Between Three and Six Months', 'Above 365 Days',
    'More than five years' ...). inf for an open-ended label, None when nothing can be read."""
    s = str(label).lower()
    for w, n in _NUMWORDS.items():
        s = re.sub(r"\b%s\b" % w, n, s)
    toks = [(float(a), u) for a, u in re.findall(r"(\d+(?:\.\d+)?)\s*(days?|d\b|months?|mos?\b|m\b|years?|yrs?|y\b)?", s)]
    if not toks:
        return None
    unit = None
    fixed = []
    for v, u in reversed(toks):                       # a unit written once applies to the numbers before it ("1-5Y")
        unit = u or unit
        fixed.append((v, unit))
    fixed.reverse()

    def months(v, u):
        u = (u or "m")[0]
        return v / 30.0 if u == "d" else (v * 12.0 if u == "y" else v)

    vals = [months(v, u) for v, u in fixed]
    if re.search(r"above|more than|over|greater|exceed|>|\+", s) and not re.search(r"less|under|below|<", s):
        return float("inf")
    return max(vals)


def _bucket_columns(buckets, sheet_cols, total_letter):
    """bucket label -> column letter, the header of each column, and a note when the two cannot be paired safely."""
    skip = re.compile(r"without|no maturity|undefined|unspecified|blank|not specified", re.I)
    cols = [(L, h) for L, h in sheet_cols if L != total_letter]
    open_cols = [(L, h) for L, h in cols if skip.search(h)]
    cols = [(L, h) for L, h in cols if not skip.search(h)]
    real = [b for b in buckets if b not in ("(blank)", "")]
    big = 1e9
    num = lambda v: big if v == float("inf") else v
    pair = {}
    his = dict((b, _upper_bound_months(b)) for b in real)
    for b in real:                                        # "1-5Y" (a range) and "5Y" (bare) both end at 60 months: the bare one is "5Y and over"
        same = [x for x in real if his[x] is not None and his[x] == his[b] and his[b] != float("inf")]
        if len(same) > 1 and not re.search(r"\d\s*(?:-|to)\s*\d", b, re.I) and any(re.search(r"\d\s*(?:-|to)\s*\d", x, re.I) for x in same):
            his[b] = float("inf")
    hi_cols = [(L, _upper_bound_months(h)) for L, h in cols]
    if all(x[1] is not None for x in hi_cols):
        for b in real:
            hb = his[b]
            if hb is None:
                continue
            best = None
            for L, hc in hi_cols:
                close = (hb == hc) or (hb != float("inf") and hc != float("inf") and abs(hb - hc) <= max(0.35, 0.12 * hc))
                if close and (best is None or abs(num(hb) - num(hc)) < best[0]):
                    best = (abs(num(hb) - num(hc)), L)
            if best:
                pair[b] = best[1]
    if len(pair) != len(real) or len(set(pair.values())) != len(real):
        pair = dict(zip(real, [L for L, _ in cols])) if len(real) == len(cols) else {}      # fall back to left-to-right order
    note = None
    if len(pair) != len(real):
        note = ("The maturity buckets of the Outstanding Report (%s) could not be paired with the maturity columns of the sheet (%s), "
                "so only the maturity total is checked." % (", ".join(real), ", ".join(h for _, h in cols)))
        pair = {}
    if "(blank)" in buckets and open_cols:
        pair["(blank)"] = open_cols[0][0]
    return pair, dict(cols + open_cols), note


def ob_templates(res, sub, notes, include_tb_ob=False):
    """Concrete rules for the off-balance sheets, from the processed Outstanding Report `res`.
    Returns (rules, pivots for the workbook)."""
    names = _ob_sheet_names(sub)
    if not any(names.values()):
        return [], None
    if res is None:
        notes.append("Off-balance sheets were extracted but no Outstanding Report is loaded, so the off-balance checks were skipped "
                     "(upload it in the Off-Balance step, then press Re-run).")
        return [], None
    try:
        from .outstanding_report import order_pivot_columns, BLANK, GRAND, FRX_COL
    except Exception:
        notes.append("The off-balance checks need the latest outstanding_report.py in the library.")
        return [], None
    # the rows behind each ACTIVE pivot (automatic, saved definition or the one the user approved): Pivot 1 feeds the main and
    # by-currency sheets, Pivot 2 the maturity sheet and the category rows
    def prep(frame):
        f = frame.copy()
        for c in ("CATEGORY", "LC_GTEE_DESC", "BILL_CCY", "BUCKET"):
            f[c] = f[c].fillna(BLANK)
        return f

    base_df = res["df"][res["df"]["_STATUS"] == "ok"]
    d1 = prep(res["data1"] if res.get("data1") is not None else base_df)
    d2 = prep(res["data2"] if res.get("data2") is not None else base_df)
    p2cols = list(res["pivot1"]["flat"]["columns"])                            # currency columns (Pivot 1)
    p1cols = list(res["pivot2"]["flat"]["columns"])                            # bucket columns (Pivot 2)
    ccys = [c for c in p2cols if c not in (GRAND, FRX_COL)]
    total_name = "Total Off-Balance Sheet Accounts"
    cats = [c for c in order_pivot_columns(list(d2["CATEGORY"].unique()), "currency") if c not in (BLANK, "(all)")]
    row_groups = load_ob_row_groups()
    grouped = {}                                                      # normalised LC/GTEE item -> the sheet row that gathers it
    for g in row_groups:
        for it in g["items"]:
            grouped[_norm_item(it)] = g["row"]

    def item_list(lc_src):
        present = list(lc_src["LC_GTEE_DESC"].unique())
        singles = [x for x in order_pivot_columns(present, "currency") if x != BLANK and _norm_item(x) not in grouped]
        gl = []
        for g in row_groups:
            have = [x for x in present if _norm_item(x) in set(_norm_item(i) for i in g["items"])]
            missing = [i for i in g["items"] if _norm_item(i) not in set(_norm_item(x) for x in present)]
            if missing and have:
                note = "'%s': not in the Outstanding Report (counted as zero): %s." % (g["row"], ", ".join(missing))
                if note not in notes:
                    notes.append(note)
            if have or missing:
                gl.append(("group", g["row"], [x for x in g["items"]]))
        return ([("total", total_name)] + [("category", c) for c in cats] + [("lc", x) for x in singles] + gl)

    items_main, items_mat = item_list(d1), item_list(d2)

    def rows_of(src, kind, name, items=None):
        if kind == "category":
            return src[src["CATEGORY"] == name]
        if kind == "lc":
            return src[src["LC_GTEE_DESC"] == name]
        if kind == "group":
            keys = set(_norm_item(i) for i in (items or []))
            return src[src["LC_GTEE_DESC"].map(_norm_item).isin(keys)]
        return src

    out = []

    def rule(label, ccy, kind, name, amount, sheet, ccy_sub, pivot, cols, letter=None, items=None):
        # `pivot` here: 2 = the currency pivot, 1 = the bucket pivot (the sheets are numbered the other way round)
        mark = None if (kind == "category" and pivot == 2) else ["ob", {2: 1, 1: 2}[pivot], kind, name, list(cols)] + ([list(items)] if items else [])
        shown = "Total" if kind == "total" else name
        side = {"match_text": name, "sheet": sheet, "currency": ccy_sub, "is_total": False, "sign": 1}
        if letter:
            side["column_letter"] = letter
        out.append({"label": "Off-balance %s - %s" % (shown, label), "currency": ccy, "rule_type": "GROUP_RULE",
                    "_expanded": True, "derived": True, "ob": True,
                    "tb_side": [{"level": "external", "label": "%s | %s" % ("Grand Total" if kind == "total" else name, label),
                                 "amount": amount, "currency": ccy, "source": OB_SOURCE, "mark": mark}],
                    "submission_side": [side]})

    # ---- main sheet and the by-currency sheet (Pivot 2)
    direct, has_iqd_col, other_col = [], False, False
    if names["foreign"]:
        sc = set(sub[sub.sheet == names["foreign"]].currency)
        has_iqd_col = "IQD" in sc
        direct = [c for c in sorted(sc) if c not in ("TOTAL", "IQD", FRX_CODE, "UNSPECIFIED") and not str(c).startswith("TOTAL_")]
        other_col = FRX_CODE in sc
        # the major currencies always get their own check, even when the sheet's header for them is not recognised: the value
        # search then looks through the unrecognised columns instead of the currency silently dropping into "Other"
        pivot_ccys = set(d1["BILL_CCY"].unique()) | set(d2["BILL_CCY"].unique())
        direct += [c for c in ("USD", "EUR", "GBP") if c in pivot_ccys and c not in direct]
    # An item gets its own checks on a sheet only when that sheet has a row of the same name (a row that merely sounds similar
    # would give a false variance). The total and the configured row groups ('For Other Purposes') are always checked.
    labels = {}
    for key in ("main", "foreign", "maturity"):
        labels[key] = sorted(set(str(x) for x in sub[sub.sheet == names[key]].line_description)) if names[key] else []
    no_row = []

    def present(kind, name, key):
        if kind in ("total", "group"):
            return True
        ok = any(_norm_item(name) == _norm_item(l) or text_similarity(name, l) >= 0.85 for l in labels[key])
        if not ok and name not in no_row:
            no_row.append(name)
        return ok

    for entry in items_main:
        kind, name = entry[0], entry[1]
        gitems = entry[2] if len(entry) > 2 else None
        rows = rows_of(d2 if kind == "category" else d1, kind, name, gitems)
        iq = float(rows[rows["BILL_CCY"] == LOCAL_CCY]["EQUI_IQD"].sum())
        fx = float(rows[rows["BILL_CCY"] != LOCAL_CCY]["EQUI_IQD"].sum())
        tt = float(rows["EQUI_IQD"].sum())
        if names["main"] and present(kind, name, "main"):
            rule("IQD (column 3 or 4)", "IQD", kind, name, iq, names["main"], "IQD", 2, [LOCAL_CCY], items=gitems)
            rule("foreign currencies (column 5 or 6)", FRX_CODE, kind, name, fx, names["main"], FRX_CODE, 2, [FRX_COL], items=gitems)
            rule("total (column 2)", "TOTAL", kind, name, tt, names["main"], "TOTAL", 2, [GRAND], items=gitems)
        if names["foreign"] and present(kind, name, "foreign"):
            if has_iqd_col:
                rule("by currency - IQD", "IQD", kind, name, iq, names["foreign"], "IQD", 2, [LOCAL_CCY], items=gitems)
            for c in direct:
                rule("by currency - %s" % c, c, kind, name, float(rows[rows["BILL_CCY"] == c]["EQUI_IQD"].sum()),
                     names["foreign"], c, 2, [c], items=gitems)
            if other_col:
                rest = rows[(rows["BILL_CCY"] != LOCAL_CCY) & (~rows["BILL_CCY"].isin(direct))]
                rule("by currency - Other Foreign Currencies", FRX_CODE, kind, name, float(rest["EQUI_IQD"].sum()), names["foreign"],
                     FRX_CODE, 2, [c for c in ccys if c != LOCAL_CCY and c not in direct], items=gitems)

    # ---- maturity sheet (Pivot 1, Bucket columns)
    if names["maturity"]:
        cols = _sheet_columns(sub, names["maturity"])
        total_letter = next((L for L, h in cols if "total" in h.lower()), cols[0][0] if cols else None)
        buckets = [b for b in p1cols if b != GRAND]
        pair, header_of, note = _bucket_columns(buckets, cols, total_letter)
        if note:
            notes.append(note)
        for entry in items_mat:
            kind, name = entry[0], entry[1]
            gitems = entry[2] if len(entry) > 2 else None
            rows = rows_of(d2, kind, name, gitems)
            if not present(kind, name, "maturity"):
                continue
            if total_letter:
                rule("maturity total", "TOTAL", kind, name, float(rows["EQUI_IQD"].sum()), names["maturity"], None, 1, [GRAND], total_letter,
                     items=gitems)
            for b, L in pair.items():
                rule("bucket %s (column %s, %s)" % (b, L, header_of.get(L, "")), "TOTAL", kind, name,
                     float(rows[rows["BUCKET"] == b]["EQUI_IQD"].sum()), names["maturity"], None, 1, [b], L, items=gitems)
    if no_row:
        notes.append("These Outstanding Report items have no row of the same name on the off-balance sheets, so they are not checked one by one "
                     "(they are inside the totals): %s. If the sheet gathers some of them under one label, add that label and its items to "
                     "\"ob_row_groups\" in default_mapping.json (like 'For Other Purposes')." % ", ".join(no_row))
    # OPTIONAL (off by default - the Trial Balance and the Outstanding Report are different sources): every OB- group of
    # the Trial Balance against the grand total of the off-balance sheets, in case the Outstanding Report is in doubt
    if include_tb_ob and names["main"]:
        for ccy, what, sub_ccy in (("IQD", "IQD (column 3 or 4)", "IQD"), (FRX_CODE, "foreign currencies (column 5 or 6)", FRX_CODE),
                                   ("TOTAL", "total (column 2)", "TOTAL")):
            out.append({"label": "Trial Balance OB- groups vs off-balance total - %s" % what, "currency": ccy, "rule_type": "GROUP_RULE",
                        "_expanded": True, "derived": True, "ob": True,
                        "tb_side": [grp("OB-*", ccy)],
                        "submission_side": [{"match_text": total_name, "sheet": names["main"], "currency": sub_ccy, "is_total": False, "sign": 1}]})
    return out, {"pivot1": res["pivot1"]["flat"], "pivot2": res["pivot2"]["flat"]}


# ------------------------------------------------------------------------------------------ the run

def _cells_text(components):
    by = OrderedDict()
    for c in components:
        if c.get("placeholder"):
            by.setdefault((c["submission_file"], c["sheet"]), []).append("(blank in this currency, row %s)" % c.get("row_number"))
        else:
            by.setdefault((c["submission_file"], c["sheet"]), []).append(c["source_cell"])
    return "; ".join("%s!%s" % (sh, "+".join(cells)) for (f, sh), cells in by.items())


VALUE_SEARCH_MIN = 10000.0          # smaller amounts are too likely to match a cell by coincidence


def _cell_comp(r, sign):
    return {"submission_file": r.submission_file, "sheet": r.sheet, "row_number": int(r.row_number), "source_cell": r.source_cell,
            "line_description": r.line_description, "currency": r.currency, "amount": float(r.normalized_amount), "multiplier": 1,
            "sign": sign, "match_confidence": 1.0, "by_value": True}


def _raw_sheet_cells(path, sheet):
    """Every numeric cell of a sheet read straight from the workbook: (cell, row number, row label, value)."""
    out = []
    try:
        ws = load_workbook(path, read_only=True, data_only=True)[sheet]
        for r, row in enumerate(ws.iter_rows(), 1):
            label = next((str(c.value).strip() for c in row if isinstance(getattr(c, "value", None), str) and len(str(c.value).strip()) > 2), "")
            for c in row:
                v = getattr(c, "value", None)
                if isinstance(v, (int, float)) and not isinstance(v, bool) and v == v:
                    out.append((c.coordinate, r, label, float(v)))
    except Exception:
        return []
    return out


_SERIAL_CACHE = {}


def _serial_rows(sources, file_label, sheet):
    """Row numbers whose first column holds a serial number (1, 2, 3 ...): the top-level lines of a CBI statement."""
    path = (sources or {}).get(file_label)
    if not path:
        return set()
    try:
        st = os.stat(path)
        key = (path, sheet, st.st_mtime_ns, st.st_size)       # a re-uploaded file under the same name is read again
    except OSError:
        return set()
    if key not in _SERIAL_CACHE:
        rows = set()
        try:
            ws = load_workbook(path, read_only=True, data_only=True)[sheet]
            for r, row in enumerate(ws.iter_rows(min_col=1, max_col=1), 1):
                v = getattr(row[0], "value", None) if row else None
                if (isinstance(v, (int, float)) and not isinstance(v, bool)) or (isinstance(v, str) and v.strip().isdigit()):
                    rows.add(r)
        except Exception:
            rows = set()
        _SERIAL_CACHE[key] = rows
        if len(_SERIAL_CACHE) > 200:
            _SERIAL_CACHE.pop(next(iter(_SERIAL_CACHE)))
    return _SERIAL_CACHE[key]


def column_value_search(sub, side, tb_amt, tol, sources=None):
    """When a rule's own row does not give the amount, look THOROUGHLY at the column that belongs to the rule's currency on the
    rule's sheet (EUR -> the EUR column of the by-currency sheet, IQD -> columns 3 and 4, foreign -> 5 and 6, Total -> the total
    column, a maturity bucket -> its column):
      1. a single cell equal to the amount (equal-and-opposite sign accepted, thousands rounding allowed);
      2. else the SUM of the column's cells: all of them, all but the rows marked total, or all but the largest cell (a
         total row that carries no 'Total' label) - a total that is not written anywhere is still found;
      3. else nothing is claimed, but the closest value is reported so the difference can be looked into.
    If the sheet has no column recognised as that currency, every column of the sheet that is not another currency's is searched.
    Returns {"kind": "cell"|"sum"|"hint", "comps": [...], ...} or None."""
    if sub is None or not len(sub) or not side.get("sheet") or abs(tb_amt) < VALUE_SEARCH_MIN:
        return None
    sheet_cells = sub[sub.sheet == side["sheet"]]
    ccy, letter, sg = side.get("currency"), side.get("column_letter"), side.get("sign", 1)
    pool, fallback = sheet_cells, False
    if ccy:
        pool = sheet_cells[sheet_cells.currency == ccy]
        if not len(pool):                       # no column was recognised as this currency: every column that is not another currency's
            pool = sheet_cells[sheet_cells.currency.isin(["UNSPECIFIED", "TOTAL", "BUCKET"]) | sheet_cells.currency.isna()]
            fallback = True
    if letter:
        pool = pool[pool.source_cell.astype(str).str.match(r"^%s\d+$" % re.escape(str(letter).upper()))]
    text = side.get("match_text") or ""
    where = side.get("currency") or "chosen"
    if fallback and sources:
        # no column of the sheet was recognised as this currency, so its cells were never extracted: read the sheet itself
        fl = sub[sub.submission_file.map(lambda x: x in sources) & (sub.sheet == side["sheet"])]
        files = sorted(set(fl.submission_file)) or sorted(sources)[:1]
        sc = float(sheet_cells.scale.iloc[0]) if len(sheet_cells) and "scale" in sheet_cells.columns else 1.0
        known = set(sheet_cells.source_cell)                       # cells that belong to a recognised (other) currency are not candidates
        best = None
        for fl_name in files:
            for cell, rown, label, val in _raw_sheet_cells(sources[fl_name], side["sheet"]):
                if cell in known:
                    continue
                v = val * sc
                for flip in (1, -1):
                    d = abs(v * sg * flip - tb_amt)
                    if d <= max(tol, 0.5 * sc if sc > 1 else 0.0):
                        key = (flip != 1, -text_similarity(text, label), d)
                        if best is None or key < best[0]:
                            best = (key, fl_name, cell, rown, label, v, flip)
        if best:
            _, fl_name, cell, rown, label, v, flip = best
            comp = {"submission_file": fl_name, "sheet": side["sheet"], "row_number": rown, "source_cell": cell, "line_description": label,
                    "currency": ccy, "amount": v, "multiplier": 1, "sign": sg * flip, "match_confidence": 1.0, "by_value": True}
            return {"kind": "cell", "comps": [comp], "n": 1, "opposite": flip == -1, "fallback": True, "where": where}
    if not len(pool):
        return None
    # 1. one cell
    cands = []
    for r in pool.itertuples(index=False):
        v = float(r.normalized_amount)
        sc = float(getattr(r, "scale", 1.0) or 1.0)
        allow = max(tol, 0.5 * sc if sc > 1 else 0.0)                 # figures stated in thousands carry +/- half a unit
        for flip in (1, -1):
            d = abs(v * sg * flip - tb_amt)
            if d <= allow:
                cands.append((flip != 1, -text_similarity(text, r.line_description), d, r, flip))
    if cands:
        if re.search(r"\btotal\b", text, re.I):          # a total sits at the bottom of its column: of several equal cells take the lowest one
            cands.sort(key=lambda c: (c[0], -int(c[3].row_number), c[2]))
        else:
            cands.sort(key=lambda c: (c[0], c[1], c[2]))
        _, _, _, r, flip = cands[0]
        return {"kind": "cell", "comps": [_cell_comp(r, sg * flip)], "n": len(cands), "opposite": flip == -1, "fallback": fallback, "where": where}
    # 2. the sum of ONE column (IQD sits in columns 3 AND 4 but a subtotal is in one of them, never their sum)
    rows = list(pool.itertuples(index=False))
    groups = {}
    for r in rows:
        m = re.match(r"^([A-Z]+)\d+$", str(r.source_cell))
        groups.setdefault((r.submission_file, m.group(1) if m else ""), []).append(r)
    for (fname, letter), rws in sorted(groups.items()):
        if len(rws) < 2:
            continue
        vals = [float(r.normalized_amount) for r in rws]
        big = max(range(len(rws)), key=lambda i: abs(vals[i]))
        serial = _serial_rows(sources, fname, side["sheet"])           # rows that carry a serial number in column A = the top-level lines
        top = [i for i, r in enumerate(rws) if int(r.row_number) in serial]
        variants = [("every cell of the column", list(range(len(rws)))),
                    ("every cell except the rows marked total", [i for i, r in enumerate(rws) if not bool(getattr(r, "is_total", False))]),
                    ("every cell except the largest one (probably the column's own total row)", [i for i in range(len(rws)) if i != big])]
        if top:                                   # a hierarchical sheet: parents already hold the sum of their children
            variants += [("the numbered top-level lines (serial no. in column A) except the largest one (the total line)", [i for i in top if i != big]),
                         ("the numbered top-level lines (serial no. in column A)", top)]
        seen = set()
        for what, idx in variants:
            key = tuple(idx)
            if len(idx) < 2 or key in seen:
                continue
            seen.add(key)
            sc_ = max(float(getattr(rws[i], "scale", 1.0) or 1.0) for i in idx)
            allow = max(tol, 0.5 * sc_ * len(idx) if sc_ > 1 else 0.0)
            for flip in (1, -1):
                if abs(sum(vals[i] for i in idx) * sg * flip - tb_amt) <= allow:
                    return {"kind": "sum", "comps": [_cell_comp(rws[i], sg * flip) for i in idx], "n": len(idx), "opposite": flip == -1,
                            "what": ("column %s: " % letter if letter else "") + what, "fallback": fallback, "where": where}
    # 3. nothing equals the amount: report the closest cell
    near = min(pool.itertuples(index=False), key=lambda r: min(abs(float(r.normalized_amount) * sg - tb_amt), abs(-float(r.normalized_amount) * sg - tb_amt)))
    v = float(near.normalized_amount) * sg
    return {"kind": "hint", "comp": _cell_comp(near, sg), "diff": tb_amt - v, "fallback": fallback, "where": where, "n_cells": len(rows)}


def ob_depth_results(res, sub, sources, eng, ob_names, notes, start_n):
    """DEPTH SEARCH of the off-balance pivots: EVERY value of Pivot 1 and of Pivot 2 is searched in every non-zero cell of the
    off-balance workbook (the bucket pivot on the maturity sheet), with the same engine as the Depth Search page (thousands /
    millions scale, equal-and-opposite sign, column-currency check). One result per distinct value, tagged "(depth search)"."""
    from .depth_search import run_depth_search
    from .outstanding_report import depth_specs
    sheets = set(x for x in ob_names.values() if x)
    labels = sorted(set(sub[sub.sheet.isin(sheets)].submission_file)) if (sub is not None and len(sub)) else []
    labels = [x for x in labels if (sources or {}).get(x)]
    if not labels:
        notes.append("The off-balance depth search needs the original workbook of the off-balance file; it was skipped.")
        return []
    import pandas as pd
    empty_tb = pd.DataFrame({"bs_mapping": [], "account": [], "account_desc": [], "tran_ccy": [], "adjusted_balance": []})
    kws = ", ".join(["maturity, tenor, tenure"] + [x.lower() for x in (ob_names.get("maturity"),) if x])
    ds = run_depth_search(empty_tb, [{"label": x, "path": sources[x]} for x in labels], dict((x, "offbalance") for x in labels),
                          {"tol_abs": eng.tolerance_abs, "maturity_keywords": kws}, externals=depth_specs(res))
    T = dict((t["id"], t) for t in ds["targets"])
    members = defaultdict(list)
    for t in ds["targets"]:
        members[t["group"]].append(t)
    from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
    out = []
    for entry in ds["files"]:
        for w in entry.get("warnings") or []:
            if w not in notes:
                notes.append(w)
    for t in ds["targets"]:
        if t["group"] != t["id"] or not t.get("external"):
            continue
        if abs(t["amount"]) < 1e-9:
            continue
        hits = []
        for entry in ds["files"]:
            hits += [h for h in entry["hits"] if h["target_id"] == t["id"]]
        hits.sort(key=lambda h: (h["currency_check"] == "CONFLICT", -float(h.get("name_similarity") or 0), abs(float(h["diff"]))))
        same = [m["label"] for m in members[t["group"]] if m["id"] != t["id"]]
        pivot_n = int(str(t["sheet"]).split()[-1])
        col_letters, row = coordinate_from_string(t["cell"])
        mark = ["obcell", pivot_n, row, column_index_from_string(col_letters)]
        warn = []
        if same:
            warn.append("The same value is also: " + "; ".join(same[:4]) + (" ..." if len(same) > 4 else "") + ".")
        comps, status, where = [], "NOT_FOUND", ""
        if hits:
            h = hits[0]
            conflict = h["currency_check"] == "CONFLICT"
            comps = [{"submission_file": h["file"], "sheet": h["sheet"], "row_number": int(h["row"]), "source_cell": h["cell"],
                      "line_description": h["row_label"], "currency": h.get("column_currency") or t["currency"],
                      "amount": float(h["value"]) * float(h["scale"]), "multiplier": 1, "sign": -1 if h["opposite_sign"] else 1,
                      "match_confidence": 1.0, "depth": True}]
            status = "REVIEW_REQUIRED" if conflict else ("MATCH" if abs(float(h["diff"])) <= 1e-6 else "MATCH_WITHIN_TOLERANCE")
            where = "%s!%s" % (h["sheet"], h["cell"])
            how = []
            if float(h["scale"]) != 1.0:
                how.append("the sheet is in x%s" % format(h["scale"], "g"))
            if h["opposite_sign"]:
                how.append("opposite sign")
            warn.insert(0, "Found in row '%s', column '%s'%s." % (h["row_label"], h["column_header"], (" (" + ", ".join(how) + ")") if how else ""))
            if conflict:
                warn.insert(0, "The value is only found in a column of a different currency (%s), so it is not accepted as a match." % h.get("column_currency"))
            if len(hits) > 1:
                warn.append("%d other cell(s) hold the same value: %s." % (len(hits) - 1, ", ".join("%s!%s" % (x["sheet"], x["cell"]) for x in hits[1:4])))
        sub_amt = float(sum(c["amount"] * c["sign"] for c in comps))
        var = float(t["amount"]) - sub_amt
        tbc = [{"account": t["label"], "account_desc": t["label"], "bs_mapping": OB_SOURCE, "currency": t["currency"], "amount": float(t["amount"]),
                "sign": 1, "external": True}]
        out.append({
            "n": start_n + len(out), "label": "Off-balance %s (depth search)" % t["label"].replace("  |  ", " | "), "rule_type": "DEPTH_SEARCH",
            "currency": t["currency"], "rule_text": "Pivot %d value searched in every non-zero cell of the off-balance sheets%s" % (
                pivot_n, " (maturity sheet only)" if t.get("scope") == "maturity" else ""),
            "tb_amount": float(t["amount"]), "sub_amount": sub_amt, "variance": var,
            "variance_pct": (var / abs(t["amount"]) * 100.0) if abs(t["amount"]) > 1e-9 else None, "status": status,
            "files": sorted(set(c["submission_file"] for c in comps)), "target_sheets": sorted(sheets), "where": where,
            "tb_components": tbc, "components": comps, "warnings": warn, "tb_marks": [mark], "expanded": True, "derived": True, "ob": True,
            "by_value": False, "depth": True, "resolved": {"tb_components": tbc, "components": comps}, "wildcard": False,
            "n_tb": 1, "n_sub": len(comps), "color": ""})
    return out


def run_default_mapping(eng, tb, sub, templates=None, outstanding=None, ob_tb=False, col_search=True, sources=None, ob_depth=True):
    templates = templates if templates is not None else eng.load_default_mapping()
    skipped = []
    concrete = expand_templates(templates, tb, sub, skipped)
    ob_notes = []
    ob_rules, ob_pivots = ob_templates(outstanding, sub, ob_notes, include_tb_ob=ob_tb)
    concrete = list(concrete) + ob_rules
    scale_of = {}
    if sub is not None and len(sub):
        for r in sub[["submission_file", "sheet", "source_cell", "scale"]].itertuples(index=False):
            scale_of[(r.submission_file, r.sheet, r.source_cell)] = float(r.scale)
    results = []
    for t in concrete:
        resolved, warnings = eng.resolve_rule_template(t, tb, sub)
        tbc = resolved.get("tb_components") or []
        sbc = resolved.get("components") or []
        tb_amt = float(sum(c["amount"] * c.get("sign", 1) for c in tbc))
        sub_amt = float(sum(c["amount"] * c.get("sign", 1) * c.get("multiplier", 1) for c in sbc))
        if t.get("_expanded") and abs(tb_amt) < 1e-9 and (not sbc or abs(sub_amt) < 1e-9):
            continue                                                  # a currency neither side has anything for
        allow = 0.0
        for c in sbc:
            sc = scale_of.get((c.get("submission_file"), c.get("sheet"), c.get("source_cell")), 1.0)
            if sc > 1:
                allow += 0.5 * sc * abs(c.get("multiplier", 1))     # figures stated in thousands carry +/- half a unit
        var = tb_amt - sub_amt
        tol = max(eng.tolerance_abs, abs(tb_amt) * eng.tolerance_pct)
        rounding = False
        if not sbc:
            status = "NOT_FOUND"
        elif abs(var) <= 1e-6:
            status = "MATCH"
        elif abs(var) <= tol:
            status = "MATCH_WITHIN_TOLERANCE"
        elif abs(var) <= allow:
            status, rounding = "MATCH_WITHIN_TOLERANCE", True
        else:
            status = "REVIEW_REQUIRED"
        notes = list(warnings)
        by_value = False
        if col_search and status in ("NOT_FOUND", "REVIEW_REQUIRED") and len(t.get("submission_side") or []) == 1:
            hit = column_value_search(sub, t["submission_side"][0], tb_amt, tol, sources)
            if hit and hit["kind"] == "hint":
                c = hit["comp"]
                notes.append("Searched the %s column(s) of '%s' (%d cells) for %s: no cell or column sum equals it. The closest value is %s in %s "
                             "(row '%s'), %s away." % (hit["where"], c["sheet"], hit["n_cells"], format(tb_amt, ",.2f"), format(c["amount"] * c["sign"], ",.2f"),
                                                       c["source_cell"], c["line_description"], format(abs(hit["diff"]), ",.2f")))
            elif hit and hit["kind"] == "sum" and sbc:
                # the sheet's own row exists and holds another number while the item rows add up to the pivot: the sheet is not
                # consistent with itself, so it stays a variance - only the explanation is added
                notes.append("The pivot amount %s equals the sum of %s (%d cells, %s) of the %s column(s) of '%s', but the sheet's own row (%s) holds %s "
                             "(%s away): the total on the sheet does not add up." % (
                                 format(tb_amt, ",.2f"), hit["what"], hit["n"], _cells_text(hit["comps"]), hit["where"], hit["comps"][0]["sheet"],
                                 _cells_text(sbc), format(sub_amt, ",.2f"), format(abs(var), ",.2f")))
            elif hit:
                comps = hit["comps"]
                was = ("its own row (%s) held %s" % (_cells_text(sbc), format(sub_amt, ",.2f"))) if sbc else "its own row was not found"
                where_txt = "searched every cell of the %s column(s) of '%s'" % (hit["where"], comps[0]["sheet"])
                if hit["fallback"]:
                    where_txt = "no %s column was recognised on '%s', so every column that is not another currency's was searched" % (
                        hit["where"], comps[0]["sheet"])
                if hit["kind"] == "cell":
                    found = "found %s in %s (row '%s')%s" % (format(tb_amt, ",.2f"), comps[0]["source_cell"], comps[0]["line_description"],
                                                            (", %d cells hold this value, %s was taken" % (hit["n"], "the lowest one (a total sits at the bottom of its column)" if re.search(r"\btotal\b", t["submission_side"][0].get("match_text") or "", re.I) else "the closest row name")) if hit["n"] > 1 else "")
                else:
                    found = "no single cell holds it, but the sum of %s (%d cells, %s) is %s" % (
                        hit["what"], hit["n"], _cells_text(comps), format(tb_amt, ",.2f"))
                notes.insert(0, "Matched by value: %s and %s; %s." % (where_txt, found, was))
                if hit["opposite"]:
                    notes.insert(1, "The value has the opposite sign.")
                sbc = comps
                sub_amt = float(sum(c["amount"] * c["sign"] for c in sbc))
                var = tb_amt - sub_amt
                allow_v = max(tol, 0.5 * sum(float(scale_of.get((c["submission_file"], c["sheet"], c["source_cell"]), 1.0)) for c in sbc))
                status = "MATCH" if abs(var) <= 1e-6 else ("MATCH_WITHIN_TOLERANCE" if abs(var) <= allow_v else "REVIEW_REQUIRED")
                resolved = dict(resolved)
                resolved["components"] = sbc
                by_value = True
        if rounding:
            notes.insert(0, "Variance is within the rounding of figures stated in thousands (+/- %s)." % format(allow, ",.0f"))
        results.append({
            "n": len(results) + 1, "label": t.get("label"), "rule_type": t.get("rule_type"), "currency": t.get("currency"),
            "rule_text": describe_rule(t), "tb_amount": tb_amt, "sub_amount": sub_amt, "variance": var,
            "variance_pct": (var / abs(tb_amt) * 100.0) if abs(tb_amt) > 1e-9 else None, "status": status,
            "files": sorted(set(c["submission_file"] for c in sbc)),
            "target_sheets": sorted(set(x.get("sheet") for x in t.get("submission_side", []) if x.get("sheet"))),
            "where": _cells_text(sbc), "tb_components": tbc, "components": sbc, "warnings": notes,
            "tb_marks": resolved.get("tb_marks", []), "expanded": bool(t.get("_expanded")), "derived": bool(t.get("derived")), "ob": bool(t.get("ob")), "by_value": by_value, "resolved": resolved,
            "wildcard": any(str(x.get("bs_mapping") or "").endswith("*") for x in (t.get("tb_side") or []) if x.get("level") == "group"),
            "n_tb": sum(1 for c in tbc if not c.get("placeholder")), "n_sub": sum(1 for c in sbc if not c.get("placeholder")),
            "color": ""})
    if ob_depth and outstanding is not None and ob_pivots:
        results += ob_depth_results(outstanding, sub, sources, eng, _ob_sheet_names(sub), ob_notes, len(results) + 1)
    palette_i = 0
    for r in results:
        if r["n_sub"]:
            r["color"] = PALETTE[palette_i % len(PALETTE)]
            palette_i += 1
    notes = ["The sheet '%s' was not extracted, so the checks that read it were skipped." % x[1] for x in skipped] + ob_notes
    return {"results": results, "summary": summarise(results), "groups": group_coverage(tb, results), "notes": notes, "ob_pivots": ob_pivots,
            "params": {"tol_abs": eng.tolerance_abs, "tol_pct": eng.tolerance_pct}}


def group_coverage(tb, results):
    """Every BS-mapping group of the Trial Balance and the rules that read it (the 'Total Assets' checks, which read
    every A- group, are not counted). A group no rule reads is listed with its amounts so it can be mapped."""
    by = defaultdict(set)
    for r in results:
        if r.get("wildcard"):
            continue
        for c in r["tb_components"]:
            if not c.get("placeholder"):
                by[str(c.get("bs_mapping")).strip()].add(r["n"])
    out = []
    if tb is None or not len(tb):
        return out
    t = tb.copy()
    t["_g"] = t.bs_mapping.astype(str).str.strip()
    for g, grp_ in t.groupby("_g", sort=True):
        local = grp_.tran_ccy == LOCAL_CCY
        is_ob = g.upper().startswith("OB")
        out.append({"group": g if g else "(blank)", "iqd": float(grp_[local].adjusted_balance.sum()),
                    "frx": float(grp_[~local].adjusted_balance.sum()), "total": float(grp_.adjusted_balance.sum()),
                    "rules": sorted(by.get(g, [])), "covered": bool(by.get(g)) or is_ob, "off_balance": is_ob and not by.get(g)})
    return out


def summarise(results):
    """Counts of the rules; the off-balance depth-search values are counted apart (they are searches, not rules)."""
    rules = [r for r in results if not r.get("depth")]
    s = {"rules": len(rules), "MATCH": 0, "MATCH_WITHIN_TOLERANCE": 0, "REVIEW_REQUIRED": 0, "NOT_FOUND": 0}
    for r in rules:
        s[r["status"]] += 1
    s["abs_variance"] = float(sum(abs(r["variance"]) for r in rules if r["status"] in ("REVIEW_REQUIRED", "MATCH_WITHIN_TOLERANCE")))
    dp = [r for r in results if r.get("depth")]
    s["depth"] = {"values": len(dp), "found": sum(1 for r in dp if r["status"] in ("MATCH", "MATCH_WITHIN_TOLERANCE")),
                  "review": sum(1 for r in dp if r["status"] == "REVIEW_REQUIRED"), "not_found": sum(1 for r in dp if r["status"] == "NOT_FOUND")}
    return s


def public_results(run):
    """JSON for the web page (no heavy internals)."""
    out = []
    for r in run["results"]:
        out.append({k: r.get(k, False if k in ("depth", "by_value", "derived", "expanded") else None) for k in ("n", "label", "rule_type", "currency", "rule_text", "tb_amount", "sub_amount", "variance",
                                      "variance_pct", "status", "files", "target_sheets", "where", "warnings", "color",
                                      "n_tb", "n_sub", "expanded", "derived", "by_value", "depth")})
    return out


# ------------------------------------------------------------------------------------------ annotated workbook

def _results_for_file(results, label, sheetnames):
    return [r for r in results if label in r["files"] or (not r["files"] and set(r["target_sheets"]) & set(sheetnames))]


def write_default_workbook(path_in, path_out, label, tb, run):
    results = run["results"]
    wb = load_workbook(path_in, keep_vba=path_in.lower().endswith(".xlsm"))
    names = set(wb.sheetnames)
    mine = _results_for_file(results, label, names)
    bold = Font(bold=True)
    grey = _fill("D9D9D9")

    # -- highlight the matched submission cells, with a comment naming the rule and its variance
    notes = defaultdict(list)
    for r in mine:
        if not r["color"]:
            continue
        for c in r["components"]:
            if c.get("submission_file") != label or c.get("placeholder") or c["sheet"] not in names:
                continue
            key = (c["sheet"], c["source_cell"])
            if not notes[key]:
                wb[c["sheet"]][c["source_cell"]].fill = _fill(r["color"])
            notes[key].append("Rule %d: %s | TB %s | variance %s" % (r["n"], r["label"][:70], format(r["tb_amount"], ",.0f"),
                                                                   format(r["variance"], ",.0f")))
    for (sheet, cell), lines in list(notes.items())[:3000]:
        cm = Comment("Default mapping:\n" + "\n".join(lines[:4]), "Default Mapping")
        cm.width, cm.height = 380, 40 + 18 * min(len(lines), 4)
        wb[sheet][cell].comment = cm

    # -- TB pivot (BS-mapping level) and TB accounts, coloured the same way
    layout = build_layout(tb, include_accounts=False, min_value=0.0)
    keys = layout["keys"]
    piv_pos = {}
    for i, (level, lab, vec, _n) in enumerate(layout["rows"]):
        if level == "BS Mapping":
            piv_pos[lab] = i + 2
    acct_pos = dict(((b, str(a), str(d)), i + 2) for i, (b, a, d, _v) in enumerate(layout["acct_rows"]))
    wp = wb.create_sheet(_unique_name(wb, "TB Pivot"))
    wp.append(["Level", "BS Mapping"] + [_col_title(k) for k in keys])
    for c in wp[1]:
        c.font, c.fill = bold, grey
    for level, lab, vec, _n in layout["rows"]:
        wp.append([level, lab] + [vec[k] for k in keys])
    wa = wb.create_sheet(_unique_name(wb, "TB Accounts"))
    wa.append(["BS Mapping", "Account", "Account Desc"] + [_col_title(k) for k in keys])
    for c in wa[1]:
        c.font, c.fill = bold, grey
    for b, a, d, vec in layout["acct_rows"]:
        wa.append([b, a, d] + [vec[k] for k in keys])
    for ws_, first in ((wp, 3), (wa, 4)):
        for row in ws_.iter_rows(min_row=2, min_col=first):
            for cell in row:
                cell.number_format = "#,##0.00;[Red]-#,##0.00"
    for r in mine:
        if not r["color"]:
            continue
        for mk in r["tb_marks"]:
            if mk[0] == "pivot" and mk[1] in piv_pos and mk[2] in keys:
                wp.cell(row=piv_pos[mk[1]], column=3 + keys.index(mk[2])).fill = _fill(r["color"])
            elif mk[0] == "acct":
                row_i = acct_pos.get((mk[1], mk[2], mk[3]))
                if row_i and mk[4] in keys:
                    wa.cell(row=row_i, column=4 + keys.index(mk[4])).fill = _fill(r["color"])
    wp.freeze_panes = "C2"
    wa.freeze_panes = "D2"
    _write_ob_pivots(wb, run, mine, bold, grey)
    wp.column_dimensions["B"].width = 36
    wa.column_dimensions["A"].width = 18
    wa.column_dimensions["C"].width = 38

    # -- the report: every rule, how it works, TB vs submission, variance
    wr = wb.create_sheet(_unique_name(wb, "Default Mapping Report"))
    wr["A1"] = "Default mapping - results for %s" % label
    wr["A1"].font = Font(bold=True, size=13)
    wr["A2"] = ("Tolerance: %s absolute or %s%% of the TB amount. Variance = TB amount - submission amount. "
                "Matched cells and the TB values they were compared with share one colour." % (
                    run["params"]["tol_abs"], format(run["params"]["tol_pct"] * 100, "g")))
    head = ["#", "Colour", "Rule", "How the rule works", "Currency", "TB amount", "Submission sheet", "Submission cells",
            "Submission amount", "Variance (TB - submission)", "Variance %", "Status", "Notes"]
    wr.append([])
    wr.append(head)
    for c in wr[4]:
        c.font, c.fill = bold, grey
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in mine:
        sheets = ", ".join(sorted(set(c["sheet"] for c in r["components"]))) or ", ".join(r["target_sheets"])
        wr.append([r["n"], "", r["label"], r["rule_text"], r["currency"], r["tb_amount"], sheets, r["where"] or "(not found)",
                   r["sub_amount"], r["variance"], r["variance_pct"], STATUS_TEXT[r["status"]], " | ".join(r["warnings"][:3])])
        ri = wr.max_row
        if r["color"]:
            wr.cell(row=ri, column=2).fill = _fill(r["color"])
        wr.cell(row=ri, column=12).fill = _fill(STATUS_HEX[r["status"]])
        for col in (6, 9, 10):
            wr.cell(row=ri, column=col).number_format = "#,##0.00;[Red]-#,##0.00"
        wr.cell(row=ri, column=11).number_format = '0.0000"%"'
        first = next((c for c in r["components"] if c.get("submission_file") == label and not c.get("placeholder")), None)
        if first and first["sheet"] in names:
            cell = wr.cell(row=ri, column=8)
            cell.hyperlink = Hyperlink(ref=cell.coordinate, location="%s!%s" % (_sheet_ref(first["sheet"]), first["source_cell"]),
                                       display=r["where"])
            cell.font = Font(color="0563C1", underline="single")
        for col in range(1, len(head) + 1):
            wr.cell(row=ri, column=col).alignment = Alignment(wrap_text=True, vertical="top")
    for letter, w in zip("ABCDEFGHIJKLM", [5, 8, 46, 70, 9, 18, 20, 30, 18, 20, 11, 22, 50]):
        wr.column_dimensions[letter].width = w
    wr.freeze_panes = "D5"
    wr.auto_filter.ref = "A4:M%d" % max(wr.max_row, 5)

    # -- every Trial Balance account behind every rule
    wd = wb.create_sheet(_unique_name(wb, "Default Mapping TB Detail"))
    wd.append(["Rule #", "Rule", "BS mapping", "Account", "Account description", "Currency", "Amount", "Currencies behind it"])
    for c in wd[1]:
        c.font, c.fill = bold, grey
    for r in mine:
        for c in r["tb_components"]:
            wd.append([r["n"], r["label"][:80], c.get("bs_mapping"), c.get("account"), c.get("account_desc"), c.get("currency"),
                       c.get("amount", 0) * c.get("sign", 1),
                       ", ".join("%s %s" % (k, format(v, ",.0f")) for k, v in (c.get("breakdown") or {}).items())])
            wd.cell(row=wd.max_row, column=7).number_format = "#,##0.00;[Red]-#,##0.00"
    for letter, w in zip("ABCDEFGH", [7, 50, 14, 12, 40, 10, 18, 40]):
        wd.column_dimensions[letter].width = w
    wd.freeze_panes = "A2"

    # -- every TB group and the rules that read it
    wg = wb.create_sheet(_unique_name(wb, "Default Mapping Groups"))
    wg.append(["BS mapping group", "IQD", "Foreign currencies", "Total", "Rules (#) that read it", "Covered?"])
    for c in wg[1]:
        c.font, c.fill = bold, grey
    for g in run.get("groups") or []:
        wg.append([g["group"], g["iqd"], g["frx"], g["total"], ", ".join(str(n) for n in g["rules"]),
                   "Yes" if g["covered"] else "NO RULE"])
        for col in (2, 3, 4):
            wg.cell(row=wg.max_row, column=col).number_format = "#,##0.00;[Red]-#,##0.00"
        if not g["covered"]:
            wg.cell(row=wg.max_row, column=6).fill = _fill(STATUS_HEX["NOT_FOUND"])
    for letter, w in zip("ABCDEF", [24, 20, 20, 20, 40, 12]):
        wg.column_dimensions[letter].width = w
    wg.freeze_panes = "A2"

    # -- summary
    ws = wb.create_sheet(_unique_name(wb, "Default Mapping Summary"))
    sm = summarise(mine)
    for a, b in [("Submission file", label), ("Generated", datetime.datetime.now().strftime("%Y-%m-%d %H:%M")),
                 ("Rules that apply to this file", sm["rules"]), ("Match", sm["MATCH"]),
                 ("Match within tolerance", sm["MATCH_WITHIN_TOLERANCE"]), ("Variance - review", sm["REVIEW_REQUIRED"]),
                 ("Not found", sm["NOT_FOUND"]), ("Total absolute variance (rules with a variance)", sm["abs_variance"])]:
        ws.append([a, b])
    for i in range(1, ws.max_row + 1):
        ws.cell(row=i, column=1).font = bold
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 40
    wb.active = wb.sheetnames.index(wr.title)
    for sh in wb.worksheets:
        sh.sheet_view.tabSelected = sh.title == wr.title
    wb.save(path_out)
    return path_out


def _write_ob_pivots(wb, run, mine, bold, grey):
    """OB Pivot 1 / OB Pivot 2 (the Outstanding Report) with the cells the off-balance rules used coloured."""
    pv = run.get("ob_pivots")
    if not pv or not any(m[0] in ("ob", "obcell") for r in mine for m in r["tb_marks"]):
        return
    sheets = {}
    for n in (1, 2):
        flat = pv["pivot%d" % n]
        ws = wb.create_sheet(_unique_name(wb, "OB Pivot %d" % n))
        nl = len(flat["label_cols"])
        ws.append(list(flat["label_cols"]) + list(flat["columns"]))
        for c in ws[1]:
            c.font, c.fill = bold, grey
        for r in flat["rows"]:
            ws.append(list(r["labels"]) + list(r["values"]) + [r["total"]])
            for col in range(nl + 1, nl + 2 + len(r["values"])):
                ws.cell(row=ws.max_row, column=col).number_format = "#,##0.00;[Red]-#,##0.00"
        ws.freeze_panes = ws.cell(row=2, column=nl + 1)
        ws.column_dimensions["A"].width = 26
        sheets[n] = (ws, flat, nl)
    for r in mine:
        if not r["color"]:
            continue
        for mk in r["tb_marks"]:
            if mk[0] == "obcell":                     # a single cell of an OB pivot (depth search)
                sheets[mk[1]][0].cell(row=mk[2], column=mk[3]).fill = _fill(r["color"])
                continue
            if mk[0] != "ob":
                continue
            _, n, kind, name, colnames = mk[:5]
            gkeys = set(_norm_item(i) for i in mk[5]) if len(mk) > 5 else set()
            ws, flat, nl = sheets[n]
            for i, row in enumerate(flat["rows"]):
                lv, lab = row["level"], row["labels"]
                hit = (kind == "total" and lv == "Grand Total") or \
                      (kind == "category" and n == 2 and lv == "Category" and lab[0] == name) or \
                      (kind == "lc" and lv == "LC/GTEE" and (lab[1] if n == 2 else lab[0]) == name) or \
                      (kind == "group" and lv == "LC/GTEE" and _norm_item(lab[1] if n == 2 else lab[0]) in gkeys)
                if not hit:
                    continue
                for cn in colnames:
                    if cn in flat["columns"]:
                        ws.cell(row=i + 2, column=nl + 1 + flat["columns"].index(cn)).fill = _fill(r["color"])


def build_zip(entries):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, path in entries:
            z.write(path, name)
    return buf.getvalue()


# ------------------------------------------------------------------------------------------ PDF

def _money(v):
    return "" if v is None else format(v, ",.2f")


def build_pdf(run, files=None, tb_name="", suggestions=None, approved=None, manual_matches=None):
    results = run["results"]
    now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
    pdf = SimplePDF(landscape=True, title="Bahrain-Iraq Recon Studio - Default Mapping", footer="Generated " + now)
    pdf.heading("Default Mapping - results", 17)
    pdf.paragraph("Trial Balance: %s    |    Tolerance: %s absolute or %s%% of the TB amount    |    Variance = TB amount - submission amount"
                  % (tb_name or "(uploaded)", run["params"]["tol_abs"], format(run["params"]["tol_pct"] * 100, "g")), 8.5)
    pdf.paragraph("Amounts are in the same units as the Trial Balance (figures a sheet states in thousands are converted). "
                  "Text outside the Latin alphabet, such as Arabic, cannot be drawn in this PDF and appears as ? - "
                  "translate the submissions first if you need it readable here.", 7.5, color=(0.4, 0.4, 0.4))
    s = run["summary"]
    pdf.heading("1. Summary", 12)
    pdf.table(["Rules applied", "Match", "Match (within tolerance)", "Variance - review", "Not found", "Total absolute variance"],
              [[s["rules"], s["MATCH"], s["MATCH_WITHIN_TOLERANCE"], s["REVIEW_REQUIRED"], s["NOT_FOUND"], _money(s["abs_variance"])]],
              [1, 1, 1.4, 1.2, 1, 1.6], size=9, aligns=["r"] * 6)
    dp = s.get("depth") or {}
    if dp.get("values"):
        pdf.paragraph("Off-balance depth search: %d value(s) of Pivot 1 and Pivot 2 were searched in every non-zero cell of the off-balance sheets - "
                      "%d found, %d found only in a column of another currency, %d not found. They are listed in section 2 with the tag (depth search)." % (
                          dp["values"], dp["found"], dp["review"], dp["not_found"]), 8.5)
    if files:
        rows = []
        for f in files:
            mine = _results_for_file(results, f["file"], f.get("sheets") or [])
            sm = summarise(mine)
            rows.append([f["file"], f.get("type_label", ""), sm["rules"], sm["MATCH"] + sm["MATCH_WITHIN_TOLERANCE"],
                         sm["REVIEW_REQUIRED"], sm["NOT_FOUND"]])
        pdf.table(["Submission file", "Type", "Rules", "Matched", "Variance - review", "Not found"], rows, [5, 2.5, 1, 1, 1.4, 1],
                  size=8, aligns=["l", "l", "r", "r", "r", "r"])

    pdf.heading("2. Results by rule", 12)
    rows, fills = [], []
    for r in results:
        pct = "" if r["variance_pct"] is None else format(r["variance_pct"], ".4f") + "%"
        rows.append([r["n"], r["label"], r["rule_text"], r["currency"], _money(r["tb_amount"]), _money(r["sub_amount"]),
                     _money(r["variance"]), pct, STATUS_TEXT[r["status"]], r["where"] or "(not found)"])
        fills.append({8: STATUS_RGB[r["status"]]})
    pdf.table(["#", "Rule", "How the rule works", "Ccy", "TB amount", "Submission amount", "Variance", "Var %", "Status", "Where in the submission"],
              rows, [0.5, 3.4, 4.6, 0.8, 1.6, 1.6, 1.5, 1.0, 1.5, 2.6], size=6.2,
              aligns=["l", "l", "l", "l", "r", "r", "r", "r", "l", "l"], fills=fills)

    groups = run.get("groups") or []
    if groups:
        pdf.heading("3. Trial Balance groups and the rules that read them", 12)
        un = [g for g in groups if not g["covered"] and (abs(g["total"]) > 1e-6)]
        pdf.paragraph("%d of %d BS-mapping groups are read by at least one rule. %s" % (
            len(groups) - len(un), len(groups),
            ("Groups no rule reads yet (they still need a mapping): " + ", ".join(g["group"] for g in un)) if un else
            "Every group with a balance is covered."), 8, color=(0.1, 0.1, 0.1))
        pdf.table(["BS mapping group", "IQD", "Foreign currencies", "Total", "Rules"],
                  [[g["group"], _money(g["iqd"]), _money(g["frx"]), _money(g["total"]),
                    ", ".join("#%d" % n for n in g["rules"]) if g["rules"] else "NO RULE"] for g in groups],
                  [2, 1.6, 1.8, 1.6, 4], size=6.6, aligns=["l", "r", "r", "r", "l"],
                  fills=[{} if g["covered"] else {4: STATUS_RGB["NOT_FOUND"]} for g in groups])
    pdf.heading("4. Rule detail - Trial Balance side and submission side", 12)
    for r in [x for x in results if not x.get("depth")]:
        pdf.ensure(70)
        pdf.paragraph("#%d  %s" % (r["n"], r["label"]), 9.5, True, gap=1)
        pdf.paragraph("Status: %s   |   TB %s   vs   submission %s   =   variance %s%s" % (
            STATUS_TEXT[r["status"]], _money(r["tb_amount"]), _money(r["sub_amount"]), _money(r["variance"]),
            "" if r["variance_pct"] is None else " (%s%%)" % format(r["variance_pct"], ".4f")), 8, color=(0.1, 0.1, 0.1), gap=1)
        pdf.paragraph("Rule: " + r["rule_text"], 7.5, color=(0.3, 0.3, 0.3), gap=3)
        tbr = []
        for c in r["tb_components"]:
            extra = ""
            if c.get("breakdown"):
                extra = "  (" + ", ".join("%s %s" % (k, format(v, ",.0f")) for k, v in c["breakdown"].items()) + ")"
            tbr.append([c.get("bs_mapping"), c.get("account"), str(c.get("account_desc")) + extra, c.get("currency"),
                        _money(c.get("amount", 0) * c.get("sign", 1))])
        tbr.append(["Total", "", "", "", _money(r["tb_amount"])])
        pdf.table(["TB - BS group", "Account", "Description", "Ccy", "Amount"], tbr, [1.3, 1, 5, 0.8, 1.6], size=6.4,
                  aligns=["l", "l", "l", "l", "r"], bold_rows=[len(tbr) - 1])
        sbr = []
        for c in r["components"]:
            sbr.append([c.get("submission_file"), c.get("sheet"), c.get("source_cell"), c.get("line_description"), c.get("currency"),
                        _money(c.get("amount")), _money(c.get("amount", 0) * c.get("sign", 1) * c.get("multiplier", 1))])
        if not sbr:
            sbr.append(["(not found)", "", "", "; ".join(r["warnings"][:2]), "", "", ""])
        sbr.append(["Total", "", "", "", "", "", _money(r["sub_amount"])])
        pdf.table(["Submission file", "Sheet", "Cell", "Line", "Ccy", "Value in sheet", "Counted"], sbr,
                  [2.2, 1.6, 0.8, 4.2, 0.7, 1.4, 1.4], size=6.4, aligns=["l", "l", "l", "l", "l", "r", "r"], bold_rows=[len(sbr) - 1])
        if r["warnings"] and r["n_sub"]:
            pdf.paragraph("Notes: " + " | ".join(r["warnings"][:4]), 6.8, color=(0.45, 0.35, 0.05), gap=6)

    if suggestions:
        pdf.heading("5. Auto-suggested matches (Mapping Studio)", 12)
        rows = []
        for i, sg in enumerate(suggestions):
            ok = bool(approved[i]) if approved and i < len(approved) else False
            rows.append([sg.get("bs_mapping"), sg.get("currency"), sg.get("rule_type"), _money(sg.get("tb_amount")),
                         _money(sg.get("suggested_submission_amount")), _money(sg.get("difference")),
                         "%d%%" % round(100 * sg.get("confidence", 0)), "Approved" if ok else "Not approved"])
        pdf.table(["BS mapping / group", "Ccy", "Rule type", "TB amount", "Submission amount", "Difference", "Confidence", "Decision"],
                  rows, [4, 0.8, 2.2, 1.6, 1.8, 1.5, 1, 1.3], size=6.4, aligns=["l", "l", "l", "r", "r", "r", "r", "l"])

    if manual_matches:
        pdf.heading("6. Manual and imported matches (Mapping Studio)", 12)
        rows = []
        for m in manual_matches:
            tb = sum(c.get("amount", 0) * c.get("sign", 1) for c in m.get("tb_components", []))
            sb = sum(c.get("amount", 0) * c.get("sign", 1) * c.get("multiplier", 1) for c in m.get("components", []))
            rows.append([m.get("label") or m.get("bs_mapping"), m.get("currency"), m.get("source", "MANUAL"),
                         len(m.get("tb_components", [])), len(m.get("components", [])), _money(tb), _money(sb), _money(tb - sb)])
        pdf.table(["Match", "Ccy", "Source", "TB items", "Submission items", "TB total", "Submission total", "Variance"], rows,
                  [5, 0.8, 1.2, 1, 1.3, 1.6, 1.8, 1.5], size=6.6, aligns=["l", "l", "l", "r", "r", "r", "r", "r"])
    return pdf.output()
