"""Depth search: look for every Trial Balance pivot value in EVERY non-zero numeric
cell of EVERY sheet of each submission workbook (schedules and matrix sheets
included, not just the sheets the structural extractor reads), then write an
annotated copy of each workbook: the original sheets, a TB Pivot sheet, a
Matching Report, and matched cells highlighted in the same colour on both sides.

Python 3.9 compatible (no match statements, no X | Y type unions).
"""
import io
import os
import re
import zipfile
from collections import defaultdict

import numpy as np
from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink

from .engine import (FORMULA_REF_PATTERN, FRX_CODE, LOCAL_CCY, ar_norm, classify_currency_header,
                     detect_scale, normalize_digits, split_hierarchy, text_similarity)

FILE_TYPES = ("assets", "liabilities", "offbalance", "any")
FILE_TYPE_LABELS = {"assets": "Assets (033)", "liabilities": "Liabilities (034)",
                    "offbalance": "Off-balance (064)", "any": "Other / search everything"}

# distinct light fills; a Trial Balance value and the cells it was found in share one
PALETTE = ["FFE699", "C6E0B4", "9DC3E6", "F4B183", "D9C2E9", "FFB3BA", "B7E1CD", "A9D0F5",
           "FAD7A0", "E8C1D8", "B5D9DC", "FFF2A8", "A8DDA0", "BFD7FF", "F6B9B9", "D8B4E2",
           "A6E3E9", "FFD3A5", "E2F0A0", "C9B6F2", "F7C6E0", "9FE2D0", "FFC9A8", "CDE0A6"]

FOREIGN_PAT = re.compile(r"foreign\s+currenc|foreign\s+exchange|\bforex\b|\bfx\b", re.I)
# the same in Arabic, written in ar_norm() form (checked against the normalised text)
FOREIGN_PAT_AR = re.compile(r"عملات?\s+(?:ال)?اجنبيه|(?:ال)?عمله\s+(?:ال)?اجنبيه|صرف\s+(?:ال)?اجنبي")
DEFAULT_PARAMS = {"min_value": 1000.0, "min_scaled": 100000.0, "tol_abs": 1.0,
                  "allow_scale": True, "allow_sign": True, "include_accounts": False,
                  "max_hits": 25,
                  # the Trial Balance's OB- groups are NOT mixed with the Outstanding Report numbers unless this is switched on
                  "include_tb_ob": False,
                  # Off-balance: the bucket (tenure) pivot is only searched on sheets whose name says maturity
                  "maturity_keywords": "maturity, tenor, tenure, استحقاق, آجال, أجل"}
SCALES_ALT = (1000.0, 1000000.0, 0.001, 0.000001)
# a sheet-wide unit note such as "(Amounts in Thousand Iraqi Dinars)" or "Equivalent in Thousand Dinars"
UNIT_NOTE = re.compile(r"\bamounts?\s+in\b|\bin\s+(thousand|million)s?\b|\bequivalent\b", re.I)
UNIT_NOTE_AR = re.compile(r"المبالغ|(?:ب)?(?:ال)?الاف|(?:ب)?(?:ال)?ملايين|(?:ال)?معادل")


def _is_unit_note(text):
    t = str(text or "")
    return bool(UNIT_NOTE.search(t) or UNIT_NOTE_AR.search(ar_norm(t)))


# ---------------------------------------------------------------- small helpers

def _numeric_string(s):
    t = normalize_digits(s).replace(",", "").replace(" ", "")      # Arabic-Indic digits and separators too
    neg = t.startswith("(") and t.endswith(")")
    if neg:
        t = t[1:-1]
    if re.fullmatch(r"-?\d+(\.\d+)?", t):
        v = float(t)
        return -v if neg else v
    return None


def _decimals(x):
    if x == int(x):
        return 0
    s = repr(x)
    if "e" in s or "E" in s:
        return 6
    return len(s.split(".")[1].rstrip("0")) if "." in s else 0


def column_currency(text):
    """Currency a column header implies. 'US Dollar' -> USD, 'Foreign Currency
    Accounts Converted to Iraqi Dinars' -> FRX (foreign, expressed in local), 'Accounts in
    Iraqi Dinars' -> IQD, 'Total' -> TOTAL, nothing recognisable -> None."""
    t = str(text or "")
    if not t.strip():
        return None
    code = classify_currency_header(t)
    if code == "LCY":
        code = LOCAL_CCY
    if code and code.startswith("TOTAL_"):
        code = code[len("TOTAL_"):]
    if FOREIGN_PAT.search(t) or FOREIGN_PAT_AR.search(ar_norm(t)):
        if code and code not in (LOCAL_CCY, "TOTAL"):
            return code
        return FRX_CODE
    return code


def currency_check(target_ccy, cell_ccy):
    """USD only against USD, IQD against IQD, Total against anything. A column we could not
    classify is UNKNOWN (kept and highlighted); a definite clash is CONFLICT (listed, not highlighted)."""
    if cell_ccy is None:
        return "UNKNOWN"
    if target_ccy == "TOTAL" or cell_ccy == "TOTAL":
        return "OK"
    if target_ccy == FRX_CODE:
        return "CONFLICT" if cell_ccy == LOCAL_CCY else "OK"
    if cell_ccy == FRX_CODE:
        return "CONFLICT" if target_ccy == LOCAL_CCY else "OK"
    return "OK" if target_ccy == cell_ccy else "CONFLICT"


def family_of(label):
    """Which kind of submission can hold this BS mapping: A-... assets, L-... liabilities, OB-... off-balance
    (and ONLY off-balance - never searched in Assets or Liabilities files)."""
    raw = str(label or "").strip()
    u = raw.upper()
    n = ar_norm(raw)
    if re.match(r"^OB([\s\-_/]|$)", u) or u.startswith("OFF") or u.startswith("CONTINGENT") or "خارج الميزانيه" in n:
        return "offbalance"
    if re.match(r"^A[\s\-_/]", u) or u.startswith("ASSET") or n.startswith("الاصول") or n.startswith("الموجودات"):
        return "assets"
    if re.match(r"^L[\s\-_/]", u) or u.startswith("LIAB") or n.startswith("المطلوبات") or n.startswith("الخصوم") \
            or n.startswith("الالتزامات"):
        return "liabilities"
    return "other"


def guess_file_type(filename, sheet_names=None):
    """033 = Assets, 034 = Liabilities, 064 = Off-balance - in English or Arabic. File name first, sheet names second."""
    def hit(text):
        t = str(text or "").lower()
        n = ar_norm(t)
        if "064" in t or "off balance" in t or "off-balance" in t or "offbalance" in t or "contingent" in t \
                or "خارج الميزانيه" in n or "خارج الميزانية" in t:
            return "offbalance"
        if "034" in t or "liabilit" in t or "capital" in t or "المطلوبات" in n or "الخصوم" in n or "الالتزامات" in n:
            return "liabilities"
        if "033" in t or "asset" in t or "الاصول" in n or "الموجودات" in n:
            return "assets"
        return None
    return hit(filename) or hit(" ".join(sheet_names or [])) or "any"


def _sheet_ref(name):
    return "'" + str(name).replace("'", "''") + "'"


# ---------------------------------------------------------------- workbook scan

class SheetIndex(object):
    def __init__(self, name, num, strs, hidden=False):
        self.name = name
        self.hidden = hidden
        self.num = num                      # (row, col, value, from_text)
        self.strs = strs                    # (row, col) -> text
        self.numpos = set((r, c) for r, c, _, _ in num)
        self.first_num_row = {}
        for r, c, _, _ in num:
            if c not in self.first_num_row or r < self.first_num_row[c]:
                self.first_num_row[c] = r
        self.by_row = defaultdict(list)
        for (r, c), s in strs.items():
            self.by_row[r].append((c, s))
        self.scale_hint = 1.0
        for (r, c) in sorted(strs):
            if r <= 15:
                sc = detect_scale(strs[(r, c)])
                if sc != 1.0:
                    self.scale_hint = sc
                    break
        self._hdr = {}

    def _text_ff(self, r, c):
        if (r, c) in self.strs:
            return self.strs[(r, c)]
        # a label visually spanning several columns sits only in its leftmost cell
        for c2 in range(c - 1, max(0, c - 5), -1):
            if (r, c2) in self.strs:
                return self.strs[(r, c2)]
            if (r, c2) in self.numpos:
                break
        return ""

    def header_parts(self, col):
        """Header strings above a column's first number, nearest first (max 3)."""
        if col in self._hdr:
            return self._hdr[col]
        fr = self.first_num_row.get(col)
        parts = []
        if fr:
            for r in range(fr - 1, max(0, fr - 8), -1):
                s = self._text_ff(r, col)
                if s and not FORMULA_REF_PATTERN.search(s):
                    parts.append(s)
                if len(parts) == 3:
                    break
        self._hdr[col] = parts
        return parts

    def header(self, col):
        """Readable header: the two nearest real header strings, top to bottom (unit notes left out)."""
        keep = [p for p in self.header_parts(col) if not _is_unit_note(p)][:2]
        return " | ".join(reversed(keep))

    def column_ccy(self, col):
        """Currency from the column's OWN header text, nearest first, so a sheet-wide '(Amounts in
        Thousand Iraqi Dinars)' note never overrides a column headed 'Euro'."""
        seen_total = False
        for p in self.header_parts(col):
            if _is_unit_note(p):
                continue
            code = column_currency(p)
            if code == "TOTAL":
                seen_total = True
            elif code:
                return code
        return "TOTAL" if seen_total else None

    def column_scale(self, col):
        for p in self.header_parts(col):
            sc = detect_scale(p)
            if sc != 1.0:
                return sc
        return self.scale_hint

    def label(self, row, col):
        best = ""
        best_left = False
        for c, s in self.by_row.get(row, []):
            if c == col or FORMULA_REF_PATTERN.search(s) or len(s) < 3:
                continue
            left = c < col
            if (left and not best_left) or (left == best_left and len(s) > len(best)):
                best, best_left = s, left
        return best[:140]


class WorkbookIndex(object):
    def __init__(self, path, label):
        self.path = path
        self.label = label
        self.sheets = []
        self.warnings = []
        wb = load_workbook(path, data_only=True)
        uncached = 0
        try:
            wf = load_workbook(path)
            for ws in wf.worksheets:
                wsv = wb[ws.title]
                for row in ws.iter_rows():
                    for cell in row:
                        if cell.data_type == "f" and wsv[cell.coordinate].value is None:
                            uncached += 1
        except Exception:
            pass
        if uncached:
            self.warnings.append(
                "%d formula cell(s) have no saved value (the file was never recalculated in Excel), so they "
                "could not be searched - open and save the file in Excel once, then upload it again." % uncached)
        rows, cols, vals, sidx, raws, texts = [], [], [], [], [], []
        for ws in wb.worksheets:
            num, strs = [], {}
            for row in ws.iter_rows():
                for cell in row:
                    v = cell.value
                    if v is None or isinstance(v, bool):
                        continue
                    if isinstance(v, (int, float)):
                        if np.isfinite(v):
                            num.append((cell.row, cell.column, float(v), False))
                    elif isinstance(v, str):
                        s = v.strip()
                        if not s:
                            continue
                        n = _numeric_string(s)
                        if n is not None:
                            num.append((cell.row, cell.column, n, True))
                        else:
                            strs[(cell.row, cell.column)] = s
            si = SheetIndex(ws.title, num, strs, hidden=(ws.sheet_state != "visible"))
            k = len(self.sheets)
            self.sheets.append(si)
            for r, c, v, ft in num:
                if v != 0:
                    rows.append(r); cols.append(c); raws.append(v); sidx.append(k); texts.append(ft)
        self.rows = np.array(rows, dtype=np.int64)
        self.cols = np.array(cols, dtype=np.int64)
        self.raw = np.array(raws, dtype=float)
        self.sidx = np.array(sidx, dtype=np.int64)
        self.from_text = np.array(texts, dtype=bool)
        absv = np.abs(self.raw)
        self.order = np.argsort(absv, kind="stable")
        self.sorted_abs = absv[self.order]

    @property
    def n_cells(self):
        return int(len(self.raw))

    def find(self, abs_target, scale, tol):
        lo = (abs_target - tol) / scale
        hi = (abs_target + tol) / scale
        i0 = np.searchsorted(self.sorted_abs, lo - 1e-9, side="left")
        i1 = np.searchsorted(self.sorted_abs, hi + 1e-9, side="right")
        return self.order[i0:i1]


# ---------------------------------------------------------------- TB pivot layout + targets

def _col_keys(tb):
    ccys = sorted(set(str(c) for c in tb["tran_ccy"].unique()), key=lambda c: (c != LOCAL_CCY, c))
    keys = list(ccys)
    if any(c != LOCAL_CCY for c in ccys):
        keys.append(FRX_CODE)
    keys.append("TOTAL")
    return ccys, keys


def _col_title(key):
    return {FRX_CODE: "FRX*", "TOTAL": "Grand Total"}.get(key, key)


def _vector(df, ccys):
    out = {}
    for c in ccys:
        out[c] = float(df.loc[df["tran_ccy"] == c, "adjusted_balance"].sum())
    out[FRX_CODE] = float(df.loc[df["tran_ccy"] != LOCAL_CCY, "adjusted_balance"].sum())
    out["TOTAL"] = float(df["adjusted_balance"].sum())
    return out


def build_layout(tb, include_accounts=False, min_value=0.0):
    """Rows/columns of the BS-mapping-level pivot (leaf mappings, group roll-ups, grand total), the
    optional account-level pivot, and one search target per non-trivial cell."""
    df = tb[["bs_mapping", "account", "account_desc", "tran_ccy", "adjusted_balance"]].copy()
    df["bs_mapping"] = df["bs_mapping"].replace("", "(Unmapped)")
    df["tran_ccy"] = df["tran_ccy"].astype(str)
    ccys, keys = _col_keys(df)

    leaves = sorted(df["bs_mapping"].unique(), key=lambda s: str(s).lower())
    descs = {}
    for lab, g in df.groupby("bs_mapping"):
        descs[lab] = list(dict.fromkeys(g["account_desc"].astype(str)))[:30]
    rows = []                                   # (level, label, vector, names)
    for lab in leaves:
        rows.append(("BS Mapping", lab, _vector(df[df["bs_mapping"] == lab], ccys), [lab] + descs[lab]))
    roll = defaultdict(list)
    for lab in leaves:
        parts = split_hierarchy(lab)
        for k in range(1, len(parts)):
            roll[" / ".join(parts[:k])].append(lab)
    for lab in sorted(roll):
        if lab in set(leaves):
            continue
        members = roll[lab]
        sub = df[df["bs_mapping"].isin(members)]
        names = [lab]
        for m in members:
            names.extend(descs[m][:5])
        rows.append(("Group", lab, _vector(sub, ccys), names[:30]))
    rows.append(("All", "Grand Total", _vector(df, ccys), []))

    targets = []

    def add_target(sheet, row_idx, col_idx, level, label, ccy_key, amount, family, names):
        # the whole-TB grand total nets assets against liabilities, so no single submission holds it
        if level == "All" or abs(amount) < max(min_value, 1e-9):
            return
        targets.append({"id": len(targets), "sheet": sheet, "cell": "%s%d" % (get_column_letter(col_idx), row_idx),
                        "level": level, "label": label, "currency": ccy_key, "amount": float(amount),
                        "family": family, "names": list(names), "currencies": [ccy_key], "scope": "all"})

    for i, (level, label, vec, names) in enumerate(rows):
        r = i + 2
        fam = None if level == "All" else family_of(label)
        for j, key in enumerate(keys):
            add_target("TB Pivot", r, 3 + j, level, label, key, vec[key], fam, names)

    acct_rows = []
    g = df.groupby(["bs_mapping", "account", "account_desc"], sort=True)
    for (bsm, acct, desc), sub in g:
        acct_rows.append((bsm, acct, desc, _vector(sub, ccys)))
    if include_accounts:
        for i, (bsm, acct, desc, vec) in enumerate(acct_rows):
            r = i + 2
            for j, key in enumerate(keys):
                add_target("TB Accounts", r, 4 + j, "Account", "%s / %s %s" % (bsm, acct, desc), key,
                           vec[key], family_of(bsm), [desc, bsm])
    return {"ccys": ccys, "keys": keys, "rows": rows, "acct_rows": acct_rows, "targets": targets}


# ---------------------------------------------------------------- the search

def group_targets(targets):
    """A group that has a single child (or a mapping held in one currency, which also equals its own
    'Grand Total' column) repeats the same figure: search it once, highlight every pivot cell carrying
    it. Same amount in a different family or sheet scope is NOT the same target."""
    by_id = dict((t["id"], t) for t in targets)
    seen = {}
    for t in targets:
        t["group"] = seen.setdefault((round(t["amount"], 2), t["family"], t.get("scope", "all"), bool(t.get("external"))), t["id"])
        if t["group"] != t["id"]:
            prim = by_id[t["group"]]
            prim["names"].extend(n for n in t["names"] if n not in prim["names"])
            if t["currency"] not in prim["currencies"]:
                prim["currencies"].append(t["currency"])


def _route(files, types, family):
    """Which files may hold a target of this family: those typed for it (plus 'search everything'
    files); with none, every other file. Off-balance (OB) values are searched ONLY in Off-balance
    files, and Off-balance files never receive Assets / Liabilities values."""
    ty = lambda f: types.get(f["label"], "any")
    if family == "offbalance":
        return [f for f in files if ty(f) in ("offbalance", "any")]
    live = [f for f in files if ty(f) != "offbalance"]
    if family in ("assets", "liabilities"):
        typed = [f for f in live if types.get(f["label"], "any") in (family, "any")]
        if any(types.get(f["label"], "any") == family for f in live):
            return typed
    return live


def run_depth_search(tb, files, file_types, params=None, externals=None):
    p = dict(DEFAULT_PARAMS)
    p.update({k: v for k, v in (params or {}).items() if v is not None})
    layout = build_layout(tb, include_accounts=bool(p["include_accounts"]), min_value=float(p["min_value"]))
    targets = layout["targets"]
    if not p.get("include_tb_ob"):
        targets = [t for t in targets if t.get("family") != "offbalance"]
        for i, t in enumerate(targets):
            t["id"] = i
    # Outstanding Report pivots (Off-balance): header row 1, data from row 2, one target per number
    for spec in externals or []:
        for et in spec["targets"]:
            if abs(et["amount"]) < max(float(p["min_value"]), 1e-9):
                continue
            targets.append({"id": len(targets), "sheet": spec["sheet"],
                            "cell": "%s%d" % (get_column_letter(et["col"]), et["row"]), "level": et["level"],
                            "label": et["label"], "currency": et["currency"], "amount": float(et["amount"]),
                            "family": "offbalance", "names": list(et["names"]), "currencies": [et["currency"]],
                            "scope": spec["scope"], "external": True})
    group_targets(targets)
    kws = [ar_norm(k.strip()) for k in str(p["maturity_keywords"]).split(",") if k.strip()]
    mat_re = re.compile("|".join(re.escape(k) for k in kws), re.I) if kws else None
    types = {f["label"]: (file_types or {}).get(f["label"]) or guess_file_type(f["label"]) for f in files}
    scales = [1.0] + (list(SCALES_ALT) if p["allow_scale"] else [])
    tol_abs = float(p["tol_abs"])
    route_labels = dict((fam, set(x["label"] for x in _route(files, types, fam)))
                        for fam in ("assets", "liabilities", "other", "offbalance"))
    # the Outstanding Report exists only to reconcile the Off-balance submission: its pivots go to files
    # the user typed Off-balance, never to "search everything" files
    route_labels["outstanding"] = set(f["label"] for f in files if types.get(f["label"]) == "offbalance")
    notes = []
    if externals and not route_labels["outstanding"]:
        notes.append("The Outstanding Report pivots were not searched: no file is typed Off-balance (064). "
                     "Set the type of the Off-balance submission under \"What is each file?\" and run again.")
    out_files = []
    for f in files:
        label = f["label"]
        entry = {"file": label, "type": types[label], "type_label": FILE_TYPE_LABELS.get(types[label], types[label]),
                 "skipped": False, "warnings": [], "hits": [], "unmatched": [], "stats": {}}
        try:
            idx = WorkbookIndex(f["path"], label)
        except Exception as exc:                        # unreadable / corrupt workbook
            entry["skipped"] = True
            entry["warnings"].append("Could not read this workbook: %s" % exc)
            out_files.append(entry)
            continue
        entry["warnings"].extend(idx.warnings)
        mat_sheets = set(k for k, sh in enumerate(idx.sheets)
                         if mat_re and (mat_re.search(sh.name) or mat_re.search(ar_norm(sh.name))))
        routed = [t for t in targets if t["group"] == t["id"] and label in route_labels["outstanding" if t.get("external") else (t["family"] or "other")]]
        skipped_maturity = [t for t in routed if t.get("scope") == "maturity" and not mat_sheets]
        if skipped_maturity:
            entry["warnings"].append(
                "No sheet in this file has a maturity/tenure name (%s), so the Outstanding Report bucket pivot "
                "(%d values) was not searched here. Adjust the keywords if the sheet is named differently."
                % (p["maturity_keywords"], len(skipped_maturity)))
            routed = [t for t in routed if t not in skipped_maturity]
        matched_targets = set()
        n_conflict = 0
        for t in routed:
            absT = abs(t["amount"])
            cands = {}
            for s in scales:
                if s > 1.0 and absT < max(float(p["min_scaled"]), 5.0 * s):
                    continue
                win = tol_abs if s <= 1.0 else max(tol_abs, 0.5 * s)
                for i in idx.find(absT, s, win):
                    if t.get("scope") == "maturity" and int(idx.sidx[i]) not in mat_sheets:
                        continue                        # bucket (tenure) values belong on the Maturity sheet only
                    v = float(idx.raw[i])
                    tol_c = tol_abs if s <= 1.0 else max(tol_abs, 0.5 * s * (10.0 ** (-_decimals(v))))
                    diff = abs(abs(v) * s - absT)
                    if diff > tol_c + 1e-9:
                        continue
                    opposite = (v > 0) != (t["amount"] > 0)
                    if opposite and not p["allow_sign"]:
                        continue
                    prev = cands.get(int(i))
                    if prev is None or (abs(np.log10(s)), diff) < (abs(np.log10(prev[0])), prev[1]):
                        cands[int(i)] = (s, diff, opposite)
            ranked = []
            for i, (s, diff, opposite) in cands.items():
                sh = idx.sheets[int(idx.sidx[i])]
                col, row = int(idx.cols[i]), int(idx.rows[i])
                header = sh.header(col)
                cell_ccy = sh.column_ccy(col)
                # a merged target may be held under several currency headings; judge it against the
                # specific ones (a 'Grand Total' heading alone would accept any column)
                specific = [c for c in t["currencies"] if c != "TOTAL"] or ["TOTAL"]
                if FRX_CODE in specific and len(specific) > 1:
                    specific.remove(FRX_CODE)   # FRX equal to one currency's amount IS that currency
                checks = [currency_check(c, cell_ccy) for c in specific]
                chk = "OK" if "OK" in checks else ("UNKNOWN" if "UNKNOWN" in checks else "CONFLICT")
                col_scale = sh.column_scale(col)
                rlabel = sh.label(row, col)
                lscore = max([text_similarity(rlabel, n) for n in t["names"]] or [0.0]) if rlabel else 0.0
                scale_pen = 0 if (s == 1.0 or s == col_scale) else 1
                ranked.append(((chk == "CONFLICT", scale_pen, round(diff, 6), -lscore), {
                    "target_id": t["id"], "file": label, "sheet": sh.name,
                    "cell": "%s%d" % (get_column_letter(col), row), "row": row, "col": col,
                    "value": float(idx.raw[i]), "matched_value": float(abs(idx.raw[i]) * s),
                    "scale": s, "opposite_sign": bool(opposite), "diff": float(diff),
                    "match_type": "EXACT" if s == 1.0 else "SCALED",
                    "scale_declared": bool(s != 1.0 and s == col_scale),
                    "row_label": rlabel, "column_header": header, "column_currency": cell_ccy or "",
                    "currency_check": chk, "name_similarity": round(float(lscore), 3),
                    "from_text": bool(idx.from_text[i])}))
            ranked.sort(key=lambda x: x[0])
            total_found = len(ranked)
            for _, h in ranked[: int(p["max_hits"])]:
                h["hits_for_target"] = total_found
                entry["hits"].append(h)
                if h["currency_check"] == "CONFLICT":
                    n_conflict += 1
            if any(h["currency_check"] != "CONFLICT" for _, h in ranked[: int(p["max_hits"])]):
                matched_targets.add(t["id"])
        # colour per matched TB value (stable order), shared by the TB pivot cell and every cell found
        colour = {}
        for t in routed:
            if t["id"] in matched_targets:
                colour[t["id"]] = PALETTE[len(colour) % len(PALETTE)]
        for h in entry["hits"]:
            h["color"] = colour.get(h["target_id"], "") if h["currency_check"] != "CONFLICT" else ""
        entry["unmatched"] = [t["id"] for t in routed if t["id"] not in matched_targets]
        entry["routed"] = [t["id"] for t in routed]
        entry["stats"] = {"sheets_scanned": len(idx.sheets), "numeric_cells": idx.n_cells,
                          "targets_searched": len(routed), "targets_matched": len(matched_targets),
                          "targets_unmatched": len(routed) - len(matched_targets),
                          "hits": len(entry["hits"]), "currency_conflicts": n_conflict,
                          "from_tb_ob": sum(1 for t in routed if t.get("family") == "offbalance" and not t.get("external")),
                          "from_outstanding": sum(1 for t in routed if t.get("external"))}
        out_files.append(entry)
    ext_store = [dict((k, v) for k, v in spec.items() if k != "targets") for spec in (externals or [])]
    return {"params": p, "types": types, "targets": targets, "files": out_files, "externals": ext_store, "notes": notes,
            "pivot_columns": [_col_title(k) for k in layout["keys"]]}


# ---------------------------------------------------------------- annotated workbook output

def _unique_name(wb, base):
    name, k = base, 2
    while name in wb.sheetnames:
        name = "%s (%d)" % (base, k)
        k += 1
    return name


def _fill(hexcolor):
    return PatternFill("solid", start_color=hexcolor, end_color=hexcolor)


def _transform_text(h):
    parts = []
    if h["scale"] != 1.0:
        parts.append(("x%g" % h["scale"]) if h["scale"] > 1 else ("/%g" % (1.0 / h["scale"])))
        if h["scale_declared"]:
            parts.append("scale stated in sheet")
    if h["opposite_sign"]:
        parts.append("opposite sign")
    return ", ".join(parts) or "exact"


def write_annotated_workbook(path_in, path_out, label, tb, result, file_entry):
    layout = build_layout(tb, include_accounts=bool(result["params"]["include_accounts"]), min_value=0.0)
    targets = {t["id"]: t for t in result["targets"]}
    keys = layout["keys"]
    keep_vba = path_in.lower().endswith(".xlsm")
    wb = load_workbook(path_in, keep_vba=keep_vba)
    note_font = Font(bold=True)

    # -- highlight matched submission cells, with a comment naming the TB value
    notes = defaultdict(list)
    n_comments = 0
    for h in file_entry["hits"]:
        if not h["color"] or h["sheet"] not in wb.sheetnames:
            continue
        ws = wb[h["sheet"]]
        cell = ws[h["cell"]]
        if not notes[(h["sheet"], h["cell"])]:
            cell.fill = _fill(h["color"])
        t = targets[h["target_id"]]
        notes[(h["sheet"], h["cell"])].append("%s | %s | TB %s  (%s)" % (
            t["label"], _col_title(t["currency"]), format(t["amount"], ",.2f"), _transform_text(h)))
    for (sheet, coord), lines in notes.items():
        if n_comments >= 3000:
            break
        c = Comment("Matches Trial Balance:\n" + "\n".join(lines[:6]), "Depth Search")
        c.width, c.height = 340, 40 + 18 * min(len(lines), 6)
        wb[sheet][coord].comment = c
        n_comments += 1

    # -- TB Pivot (BS-mapping level), same colours on the matched TB cells
    found_at = defaultdict(list)
    for h in file_entry["hits"]:
        if h["color"]:
            found_at[h["target_id"]].append("%s!%s" % (h["sheet"], h["cell"]))
    colour_of = {}
    for h in file_entry["hits"]:
        if h["color"]:
            colour_of[h["target_id"]] = h["color"]
    tb_by_cell = defaultdict(dict)
    members = defaultdict(list)
    for t in result["targets"]:
        tb_by_cell[t["sheet"]][t["cell"]] = t
        members[t["group"]].append(t)

    def also_at(tid):
        return "; ".join("%s: %s (%s)" % (m["level"], m["label"], _col_title(m["currency"]))
                         for m in members[tid] if m["id"] != tid)

    piv_name = _unique_name(wb, "TB Pivot")
    wp = wb.create_sheet(piv_name)
    wp.append(["Level", "BS Mapping"] + [_col_title(k) for k in keys])
    for c in wp[1]:
        c.font = note_font; c.fill = _fill("D9D9D9")
    for level, lab, vec, _names in layout["rows"]:
        wp.append([level, lab] + [vec[k] for k in keys])
    for row in wp.iter_rows(min_row=2, min_col=3):
        for cell in row:
            cell.number_format = "#,##0.00;[Red]-#,##0.00"
            t = tb_by_cell["TB Pivot"].get(cell.coordinate)
            if t and t["group"] in colour_of:
                cell.fill = _fill(colour_of[t["group"]])
                c = Comment("Found in:\n" + "\n".join(found_at[t["group"]][:8]), "Depth Search")
                c.width, c.height = 300, 40 + 18 * min(len(found_at[t["group"]]), 8)
                cell.comment = c
    wp.column_dimensions["A"].width = 12
    wp.column_dimensions["B"].width = 38
    for j in range(len(keys)):
        wp.column_dimensions[get_column_letter(3 + j)].width = 20
    wp.freeze_panes = "C2"

    acct_name = _unique_name(wb, "TB Accounts")
    wa = wb.create_sheet(acct_name)
    wa.append(["BS Mapping", "Account", "Account Desc"] + [_col_title(k) for k in keys])
    for c in wa[1]:
        c.font = note_font; c.fill = _fill("D9D9D9")
    for bsm, acct, desc, vec in layout["acct_rows"]:
        wa.append([bsm, acct, desc] + [vec[k] for k in keys])
    for row in wa.iter_rows(min_row=2, min_col=4):
        for cell in row:
            cell.number_format = "#,##0.00;[Red]-#,##0.00"
            t = tb_by_cell["TB Accounts"].get(cell.coordinate)
            if t and t["group"] in colour_of:
                cell.fill = _fill(colour_of[t["group"]])
    wa.column_dimensions["A"].width = 24
    wa.column_dimensions["C"].width = 36
    wa.freeze_panes = "D2"

    # -- Outstanding Report pivots (Off-balance): the same colours on the matched pivot cells
    if file_entry["type"] in ("offbalance", "any"):
        for spec in result.get("externals") or []:
            we = wb.create_sheet(_unique_name(wb, spec["sheet"]))
            we.append(list(spec["headers"]))
            for c in we[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = _fill("1F3864")
            for row in spec["rows"]:
                we.append(list(row))
            for rr in spec["bold"]:
                for c in we[rr]:
                    c.font = Font(bold=True)
            nl = spec["n_label"]
            for row in we.iter_rows(min_row=2, min_col=nl + 1):
                for cell in row:
                    cell.number_format = '#,##0;[Red](#,##0);"-"'
                    t = tb_by_cell[spec["sheet"]].get(cell.coordinate)
                    if t and t["group"] in colour_of:
                        cell.fill = _fill(colour_of[t["group"]])
                        c = Comment("Found in:\n" + "\n".join(found_at[t["group"]][:8]), "Depth Search")
                        c.width, c.height = 300, 40 + 18 * min(len(found_at[t["group"]]), 8)
                        cell.comment = c
            we.freeze_panes = we.cell(row=2, column=nl + 1)
            for j in range(1, len(spec["headers"]) + 1):
                we.column_dimensions[get_column_letter(j)].width = 30 if j <= nl else 18

    # -- Matching Report: what matches with what
    rep_name = _unique_name(wb, "Matching Report")
    wr = wb.create_sheet(rep_name)
    head = ["#", "Colour", "TB level", "TB BS mapping / item", "TB currency", "TB amount",
            "Submission file", "Sheet", "Cell", "Cell value (as in sheet)", "Transform", "Value after transform",
            "Difference", "Match type", "Row label", "Column header", "Column currency", "Currency check",
            "Name similarity", "TB pivot cell", "Cells holding this amount", "Same TB amount also at"]
    wr.append(head)
    for c in wr[1]:
        c.font = note_font; c.fill = _fill("D9D9D9"); c.alignment = Alignment(wrap_text=True, vertical="top")
    order = sorted(file_entry["hits"], key=lambda h: (h["target_id"], h["currency_check"] == "CONFLICT", h["diff"]))
    for n, h in enumerate(order, 1):
        t = targets[h["target_id"]]
        wr.append([n, "", t["level"], t["label"], _col_title(t["currency"]), t["amount"], label, h["sheet"], h["cell"],
                   h["value"], _transform_text(h), h["value"] * h["scale"] * (-1 if h["opposite_sign"] else 1),
                   h["diff"], h["match_type"], h["row_label"], h["column_header"], h["column_currency"],
                   h["currency_check"], h["name_similarity"], "%s!%s" % (t["sheet"], t["cell"]), h["hits_for_target"],
                   also_at(t["id"])])
        r = wr.max_row
        if h["color"]:
            wr.cell(row=r, column=2).fill = _fill(h["color"])
        cell = wr.cell(row=r, column=9)
        cell.hyperlink = Hyperlink(ref=cell.coordinate, location="%s!%s" % (_sheet_ref(h["sheet"]), h["cell"]), display=h["cell"])
        cell.font = Font(color="0563C1", underline="single")
        wr.cell(row=r, column=6).number_format = "#,##0.00;[Red]-#,##0.00"
        wr.cell(row=r, column=10).number_format = "#,##0.######;[Red]-#,##0.######"
        wr.cell(row=r, column=12).number_format = "#,##0.00"
        if h["currency_check"] == "CONFLICT":
            for cc in range(1, len(head) + 1):
                wr.cell(row=r, column=cc).font = Font(color="808080", italic=True)
    for letter, w in zip("ABCDEFGHIJKLMNOPQRSTUV", [5, 8, 11, 34, 10, 20, 26, 26, 9, 20, 22, 20, 12, 11, 40, 40, 10, 11, 10, 18, 12, 36]):
        wr.column_dimensions[letter].width = w
    wr.freeze_panes = "A2"
    wr.auto_filter.ref = wr.dimensions

    # -- Unmatched TB values for this file
    un_name = _unique_name(wb, "Unmatched TB Values")
    wu = wb.create_sheet(un_name)
    wu.append(["TB level", "TB BS mapping / item", "TB currency", "TB amount", "TB pivot cell", "Same TB amount also at"])
    for c in wu[1]:
        c.font = note_font; c.fill = _fill("D9D9D9")
    for tid in file_entry["unmatched"]:
        t = targets[tid]
        wu.append([t["level"], t["label"], _col_title(t["currency"]), t["amount"], "%s!%s" % (t["sheet"], t["cell"]), also_at(tid)])
        wu.cell(row=wu.max_row, column=4).number_format = "#,##0.00;[Red]-#,##0.00"
    wu.column_dimensions["B"].width = 44
    wu.column_dimensions["D"].width = 20
    wu.column_dimensions["E"].width = 18

    # -- Summary / legend
    sm_name = _unique_name(wb, "Depth Search Summary")
    wsm = wb.create_sheet(sm_name)
    st = file_entry["stats"]
    p = result["params"]
    rows = [("Submission file", label), ("File type", file_entry["type_label"]),
            ("Sheets scanned (all)", st.get("sheets_scanned")), ("Non-zero numeric cells scanned", st.get("numeric_cells")),
            ("TB values searched", st.get("targets_searched")), ("TB values found", st.get("targets_matched")),
            ("TB values not found", st.get("targets_unmatched")), ("Hits listed", st.get("hits")),
            ("Hits with a currency conflict (listed, not highlighted)", st.get("currency_conflicts")), ("", ""),
            ("Minimum TB value searched", p["min_value"]), ("Minimum TB value for a x1000 / x1,000,000 match", p["min_scaled"]),
            ("Tolerance (absolute)", p["tol_abs"]), ("Thousands / millions scale tried", p["allow_scale"]),
            ("Opposite sign accepted", p["allow_sign"]), ("Account-level values searched", p["include_accounts"]), ("", ""),
            ("How to read", "A TB value and every sheet cell where it was found share one fill colour. Hover a coloured cell for the comment. "
                            "Rows in grey italics in the Matching Report are currency conflicts (e.g. an IQD amount found in a USD column).")]
    for w in file_entry["warnings"]:
        rows.append(("Warning", w))
    for a, b in rows:
        wsm.append([a, b])
    for r in range(1, wsm.max_row + 1):
        wsm.cell(row=r, column=1).font = note_font
    wsm.column_dimensions["A"].width = 52
    wsm.column_dimensions["B"].width = 110
    wb.active = wb.sheetnames.index(rep_name)
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = (ws.title == rep_name)
    wb.save(path_out)
    return path_out


def build_zip(entries):
    """entries: list of (filename, path)"""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, path in entries:
            z.write(path, name)
    return buf.getvalue()
