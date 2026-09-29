import io, os, re, json, difflib, itertools, tempfile
from datetime import datetime
from pathlib import Path
from collections import defaultdict
import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter, column_index_from_string

ALIASES = {
 "account": ["account", "gl account", "account no", "account number"],
 "account_desc": ["account desc", "account description", "gl description"],
 "tran_ccy": ["tran ccy", "transaction currency", "currency", "ccy"],
 "bs_mapping": ["bs mapping", "balance sheet mapping", "bs group", "mapping"],
 "adjusted_balance": ["adjusted balance", "adjusted bal", "ytd adjusted balance"],
 "base_amt": ["base amt", "base amount", "local currency amount", "lc amount"],
 "adjustment": ["adjustment", "adjustments"],
 "category": ["cat", "category"],
 "pl_mapping": ["p&l mapping", "p/l mapping", "pl mapping"]
}

CURRENCY_CODES = ["USD","IQD","BHD","EUR","GBP","AED","SAR","KWD","QAR","OMR","JOD","CHF","JPY","CNY"]
CURRENCY_PATTERN = re.compile(r"\b(" + "|".join(CURRENCY_CODES) + r")\b", re.I)
LOCAL_CCY_PATTERN = re.compile(r"\blocal\s*currency\b|\blcy\b", re.I)
TOTAL_PATTERN = re.compile(r"\b(grand\s+total|sub\s*-?\s*total|total)\b", re.I)
TOTAL_PREFIX = re.compile(r"^(grand\s+total\s+of\s+|grand\s+total\s+|sub[\s-]?total\s+of\s+|sub[\s-]?total\s+|total\s+of\s+|total\s+)", re.I)
SCALE_THOUSANDS = re.compile(r"\(?'?000\b\)?|in\s+thousands|thousands", re.I)
SCALE_MILLIONS = re.compile(r"\(?'?mn'?\)?|in\s+millions|millions", re.I)
HIERARCHY_DELIMS = [" > ", " / ", " - ", " – ", " :: ", " | "]
STOPWORDS = {"and","the","of","for","in","on","to","a","an","with","&"}

def strip_total_prefix(text):
    t = str(text or "").strip()
    s = TOTAL_PREFIX.sub("", t).strip()
    return s or t

def detect_scale(text):
    t = str(text or "")
    if SCALE_THOUSANDS.search(t): return 1000.0
    if SCALE_MILLIONS.search(t): return 1000000.0
    return 1.0

def norm(v):
    return re.sub(r"[^a-z0-9]+", " ", str(v or "").strip().lower()).strip()

def norm_tokens(v):
    t = re.sub(r"[^a-z0-9\s]", " ", str(v or "").lower())
    return [w for w in t.split() if w and w not in STOPWORDS]

def text_similarity(a, b):
    ta, tb = set(norm_tokens(a)), set(norm_tokens(b))
    seq = difflib.SequenceMatcher(None, str(a or "").lower(), str(b or "").lower()).ratio()
    if not ta or not tb:
        return seq
    jaccard = len(ta & tb) / len(ta | tb) if (ta | tb) else 0.0
    return max(jaccard, seq)

def classify_currency_header(text):
    t = str(text or "").strip()
    if not t:
        return None
    is_total = bool(TOTAL_PATTERN.search(t))
    m = CURRENCY_PATTERN.search(t)
    code = m.group(1).upper() if m else ("LCY" if LOCAL_CCY_PATTERN.search(t) else None)
    if is_total and code:
        return f"TOTAL_{code}"
    if is_total:
        return "TOTAL"
    return code

def split_hierarchy(text):
    t = str(text or "").strip()
    if not t:
        return []
    for d in HIERARCHY_DELIMS:
        if d in t:
            return [p.strip() for p in t.split(d) if p.strip()]
    return [t]

def json_value(v):
    if pd.isna(v): return None
    if isinstance(v, (np.integer,)): return int(v)
    if isinstance(v, (np.floating,)): return float(v)
    if isinstance(v, (pd.Timestamp, datetime)): return v.isoformat()
    return str(v) if not isinstance(v,(str,int,float,bool)) else v

class ReconEngine:
    def __init__(self, tolerance_abs=1.0, tolerance_pct=0.0001):
        self.tolerance_abs=float(tolerance_abs); self.tolerance_pct=float(tolerance_pct)

    # ---------- workbook inspection ----------

    def _read_raw(self, path, sheet):
        return pd.read_excel(path, sheet_name=sheet, header=None, dtype=object, engine='openpyxl')

    def inspect_workbook(self, path):
        xls=pd.ExcelFile(path, engine='openpyxl'); out=[]
        for sheet in xls.sheet_names:
            raw=self._read_raw(path,sheet)
            header,score=self.detect_header(raw)
            headers=[] if header is None else [str(x).strip() if pd.notna(x) else '' for x in raw.iloc[header].tolist()]
            currency_columns={}
            for j,h in enumerate(headers):
                code=classify_currency_header(h)
                if code: currency_columns[get_column_letter(j+1)]={'currency':code,'scale':detect_scale(h)}
            out.append({'sheet':sheet,'rows':int(raw.shape[0]),'columns':int(raw.shape[1]),
                        'header_row':None if header is None else int(header+1),'score':score,
                        'headers':headers[:60],'currency_columns':currency_columns})
        return out

    def detect_header(self, raw, max_rows=80):
        best=(None,-1)
        vocabulary=set(sum(ALIASES.values(),[])) | {'amount','description','line item','particulars','total'}
        for i in range(min(max_rows,len(raw))):
            cells=[norm(x) for x in raw.iloc[i].tolist() if pd.notna(x)]
            score=sum(3 for c in cells if c in vocabulary)+sum(1 for c in cells if any(k in c for k in ['account','amount','balance','mapping','description']))
            if len(set(cells))>=3 and score>best[1]: best=(i,score)
        return best if best[1]>=3 else (None,best[1])

    def _canonicalize(self, df):
        cols={norm(c):c for c in df.columns}; rename={}
        for target, aliases in ALIASES.items():
            for a in aliases:
                if norm(a) in cols: rename[cols[norm(a)]]=target; break
        return df.rename(columns=rename), rename

    # ---------- Trial Balance ----------

    def load_tb(self, path, sheet=None, header_row=None):
        inspections=self.inspect_workbook(path)
        if not inspections: raise ValueError('No worksheets found.')
        pick=next((x for x in inspections if x['sheet']==sheet),None) if sheet else max(inspections,key=lambda x:x['score'])
        if not pick: raise ValueError('Selected worksheet was not found.')
        hr=int(header_row)-1 if header_row else pick['header_row']-1 if pick['header_row'] else None
        if hr is None: raise ValueError('Could not detect the Trial Balance header row.')
        df=pd.read_excel(path,sheet_name=pick['sheet'],header=hr,dtype=object,engine='openpyxl')
        df.columns=[str(c).strip() for c in df.columns]
        df,rename=self._canonicalize(df); required=['account','account_desc','tran_ccy','bs_mapping','adjusted_balance']
        missing=[x for x in required if x not in df.columns]
        if missing: raise ValueError('Missing required Trial Balance fields: '+', '.join(missing))
        for c in ['account','account_desc','tran_ccy','bs_mapping']:
            df[c]=df[c].fillna('').astype(str).str.strip()
        df['adjusted_balance']=pd.to_numeric(df['adjusted_balance'],errors='coerce').fillna(0.0)
        for c in ['base_amt','adjustment']:
            if c in df: df[c]=pd.to_numeric(df[c],errors='coerce').fillna(0.0)
        df=df.loc[~((df['account']=='') & (df['account_desc']=='') & (df['adjusted_balance']==0))].copy()
        df['_source_row']=np.arange(hr+2,hr+2+len(df)); df['_source_sheet']=pick['sheet']
        return df, {'sheet':pick['sheet'],'header_row':hr+1,'rows':len(df),'column_map':{str(k):v for k,v in rename.items()}}

    def make_pivot(self, tb):
        p=pd.pivot_table(tb,index=['bs_mapping','account','account_desc'],columns='tran_ccy',values='adjusted_balance',aggfunc='sum',fill_value=0,margins=True,margins_name='Grand Total').reset_index()
        p.columns=[str(c) for c in p.columns]
        return p

    def build_tb_tree(self, tb):
        tb=tb.copy()
        tb['_hierarchy']=tb['bs_mapping'].apply(split_hierarchy)
        tb['_top']=tb['_hierarchy'].apply(lambda h: h[0] if h else '(Unmapped)')
        tb['_sub']=tb['_hierarchy'].apply(lambda h: h[1] if len(h)>1 else '')
        tree=[]
        for top, gdf in tb.groupby('_top', sort=False):
            node={'name':top,'level':'group','currency_totals':self._currency_totals(gdf),'children':[]}
            subs=gdf[gdf['_sub']!='']
            if len(subs):
                for sub,sdf in subs.groupby('_sub', sort=False):
                    node['children'].append({'name':sub,'level':'subgroup','currency_totals':self._currency_totals(sdf),'accounts':self._account_rows(sdf),'children':[]})
                leftover=gdf[gdf['_sub']=='']
                if len(leftover):
                    node['children'].append({'name':'(Direct)','level':'subgroup','currency_totals':self._currency_totals(leftover),'accounts':self._account_rows(leftover),'children':[]})
            else:
                node['accounts']=self._account_rows(gdf)
            tree.append(node)
        return tree

    def _currency_totals(self, df):
        out={str(k):float(v) for k,v in df.groupby('tran_ccy')['adjusted_balance'].sum().to_dict().items()}
        out['TOTAL']=float(df['adjusted_balance'].sum())
        return out

    def _account_rows(self, df):
        return [{'account':r.account,'account_desc':r.account_desc,'currency':r.tran_ccy,'amount':float(r.adjusted_balance),'bs_mapping':r.bs_mapping} for r in df.itertuples()]

    # ---------- Submissions ----------

    def _to_number(self, v):
        if isinstance(v,(int,float,np.number)) and not pd.isna(v): return float(v)
        if isinstance(v,str):
            s=v.replace(',','').replace('(','-').replace(')','').replace('%','').strip()
            if re.fullmatch(r'-?\d+(\.\d+)?', s) and s not in ('','-'): return float(s)
        return None

    def extract_submission(self, path, file_label=None, sheets=None):
        file_label = file_label or Path(path).name
        inspections = self.inspect_workbook(path)
        insp_by_name = {x['sheet']: x for x in inspections}
        if sheets:
            targets=[]
            for s in sheets:
                name = s['sheet'] if isinstance(s, dict) else s
                info = dict(insp_by_name.get(name) or {'sheet':name,'header_row':None,'currency_columns':{}})
                override = s.get('header_row') if isinstance(s, dict) else None
                if override:
                    info['header_row']=int(override)
                    raw=self._read_raw(path,name); hr=int(override)-1
                    headers=[str(x).strip() if pd.notna(x) else '' for x in raw.iloc[hr].tolist()] if hr < len(raw) else []
                    info['currency_columns']={get_column_letter(j+1):{'currency':classify_currency_header(h),'scale':detect_scale(h)} for j,h in enumerate(headers) if classify_currency_header(h)}
                targets.append(info)
        else:
            targets = inspections

        wb=load_workbook(path, data_only=True)
        rows=[]
        for info in targets:
            sheet_name=info['sheet']
            if sheet_name not in wb.sheetnames: continue
            raw=self._read_raw(path, sheet_name); ws=wb[sheet_name]
            header_row = info['header_row']-1 if info.get('header_row') else None
            currency_cols={column_index_from_string(k)-1:v for k,v in (info.get('currency_columns') or {}).items()}
            has_any_currency=bool(currency_cols)
            stack=[]
            for r in range(len(raw)):
                if header_row is not None and r<=header_row: continue
                vals=raw.iloc[r].tolist()
                label_col,label=None,''
                for j,v in enumerate(vals):
                    if pd.notna(v) and str(v).strip():
                        label_col,label=j,str(v).strip(); break
                if label_col is None: continue
                numeric_cells={}
                for j,v in enumerate(vals):
                    if j==label_col: continue
                    num=self._to_number(v)
                    if num is not None: numeric_cells[j]=num
                cell=ws.cell(row=r+1, column=label_col+1)
                indent=0
                try: indent=int(cell.alignment.indent or 0)
                except Exception: indent=0
                bold=bool(cell.font and cell.font.bold)
                is_total=bool(TOTAL_PATTERN.search(label))
                is_header_row = not numeric_cells
                if is_header_row:
                    if bold or label.isupper() or indent>0:
                        level = indent if indent else 0
                        while stack and stack[-1][0] >= level: stack.pop()
                        stack.append((level, label))
                    continue
                hierarchy_path=[name for _,name in stack]
                level = indent if indent else len(stack)
                for j, amount in numeric_cells.items():
                    colinfo=currency_cols.get(j)
                    if not colinfo:
                        if has_any_currency: continue
                        ccy,scale='UNSPECIFIED',1.0
                    else:
                        ccy,scale=colinfo.get('currency','UNSPECIFIED'),float(colinfo.get('scale',1.0))
                    rows.append({
                        'submission_file':file_label,'sheet':sheet_name,'row_number':r+1,
                        'source_cell':f'{get_column_letter(j+1)}{r+1}','label_cell':f'{get_column_letter(label_col+1)}{r+1}',
                        'line_description':label,'hierarchy':hierarchy_path,
                        'hierarchy_path':' / '.join(hierarchy_path) if hierarchy_path else '',
                        'level':level,'is_total':is_total,'currency':ccy,'scale':scale,
                        'raw_amount':amount,'normalized_amount':amount*scale,'unit_multiplier':scale,'sign':1.0
                    })
        return pd.DataFrame(rows)

    # ---------- Auto-matching ----------

    def auto_match(self, tb_tree, submission_df, min_group_score=0.35, min_sub_score=0.3):
        if submission_df is None or submission_df.empty: return []
        recs=submission_df.to_dict('records')
        sub_top=defaultdict(list)
        for rec in recs:
            hier = rec.get('hierarchy') or []
            top = hier[0] if hier else rec['line_description']
            sub_top[top].append(rec)

        suggestions=[]
        for group in tb_tree:
            gname=group['name']
            best_top,best_score=None,0.0
            for top in sub_top:
                score=text_similarity(gname, top)
                if score>best_score: best_top,best_score=top,score
            if best_score < min_group_score or best_top is None: continue
            top_rows=sub_top[best_top]
            children = group.get('children') or [{'name':'(Direct)','level':'subgroup','currency_totals':group.get('currency_totals',{}),'accounts':group.get('accounts',[])}]
            for child in children:
                cname=child['name']
                sub_candidates=defaultdict(list)
                for r in top_rows:
                    hier=r.get('hierarchy') or []
                    key = hier[1] if len(hier)>1 else strip_total_prefix(r['line_description'])
                    sub_candidates[key].append(r)
                best_sub,best_sub_score=None,0.0
                for key in sub_candidates:
                    score=text_similarity(cname,key)
                    if score>best_sub_score: best_sub,best_sub_score=key,score
                pool = sub_candidates.get(best_sub, top_rows) if (best_sub_score>=min_sub_score and cname!='(Direct)') else top_rows
                for ccy, tb_amt in (child.get('currency_totals') or {}).items():
                    if ccy=='TOTAL':
                        matches=[r for r in pool if r['currency']=='TOTAL' or str(r['currency']).startswith('TOTAL_')]
                    else:
                        matches=[r for r in pool if r['currency']==ccy]
                    if not matches: continue
                    totals=[r for r in matches if r['is_total']]
                    chosen = totals if totals else matches
                    raw_sum=sum(r['normalized_amount'] for r in chosen)
                    resolved_sign = -1.0 if abs(tb_amt-(-raw_sum)) < abs(tb_amt-raw_sum) else 1.0
                    sub_amt=raw_sum*resolved_sign
                    diff=tb_amt-sub_amt
                    label = gname if cname in ('(Direct)','') else f'{gname} / {cname}'
                    suggestions.append({
                        'bs_mapping':label,'group':gname,'subgroup':None if cname=='(Direct)' else cname,
                        'currency':ccy,'tb_amount':tb_amt,'suggested_submission_amount':sub_amt,'difference':diff,
                        'sign_applied':resolved_sign,
                        'confidence':round(min(best_score, best_sub_score if best_sub_score else best_score),2),
                        'match_basis':{'section_match':best_top,'subgroup_match':best_sub,'used_total_rows':bool(totals)},
                        'rule_type':'AUTO_TOTAL' if ccy=='TOTAL' else 'AUTO_CURRENCY',
                        'components':[{'submission_file':r['submission_file'],'sheet':r['sheet'],'row_number':r['row_number'],
                                       'source_cell':r['source_cell'],'line_description':r['line_description'],'currency':r['currency'],
                                       'amount':r['normalized_amount'],'multiplier':1,'sign':resolved_sign} for r in chosen]
                    })
        suggestions.sort(key=lambda s: (-s['confidence'], s['bs_mapping']))
        return suggestions

    # ---------- Reconciliation ----------

    def reconcile(self, tb, submissions, rules):
        tb=tb.copy()
        tb['_hier_key']=tb['bs_mapping'].apply(lambda x: ' / '.join(split_hierarchy(x)))
        grp_raw_total=tb.groupby('bs_mapping')['adjusted_balance'].sum().to_dict()
        grp_raw_ccy=tb.groupby(['bs_mapping','tran_ccy'])['adjusted_balance'].sum().to_dict()
        grp_hier_total=tb.groupby('_hier_key')['adjusted_balance'].sum().to_dict()
        grp_hier_ccy=tb.groupby(['_hier_key','tran_ccy'])['adjusted_balance'].sum().to_dict()
        out=[]
        for rule in rules:
            bs=rule.get('bs_mapping','').strip(); ccy=rule.get('currency','TOTAL') or 'TOTAL'
            comps=rule.get('components',[]); sub_amt=0.0; evidence=[]
            for comp in comps:
                f,s,row=comp.get('submission_file'),comp.get('sheet'),comp.get('row_number')
                if row in (None,''): continue
                cell=comp.get('source_cell'); comp_ccy=comp.get('currency')
                q=(submissions.submission_file==f)&(submissions.sheet==s)&(submissions.row_number==int(row))
                if cell: q=q&(submissions.source_cell==cell)
                elif comp_ccy: q=q&(submissions.currency==comp_ccy)
                hit=submissions[q]
                if len(hit):
                    r0=hit.iloc[0]
                    amount=float(r0.normalized_amount)*float(comp.get('multiplier',1))*float(comp.get('sign',1))
                    sub_amt+=amount
                    evidence.append({'file':f,'sheet':s,'row':int(row),'description':r0.line_description,'amount':amount,'currency':str(r0.get('currency',''))})
            if ccy=='TOTAL':
                tb_amt=float(grp_raw_total.get(bs, grp_hier_total.get(bs,0)))
            else:
                tb_amt=float(grp_raw_ccy.get((bs,ccy), grp_hier_ccy.get((bs,ccy),0)))
            diff=tb_amt-sub_amt; pct=abs(diff)/max(abs(tb_amt),1)
            status='MATCH' if abs(diff)<1e-9 else 'MATCH_WITHIN_TOLERANCE' if abs(diff)<=self.tolerance_abs or pct<=self.tolerance_pct else 'REVIEW_REQUIRED'
            out.append({'bs_mapping':bs,'currency':ccy,'tb_amount':tb_amt,'submission_amount':sub_amt,'difference':diff,
                        'variance_pct':pct,'status':status,'rule_type':rule.get('rule_type','DIRECT'),
                        'source':rule.get('source','MANUAL'),'evidence':json.dumps(evidence,ensure_ascii=False)})
        return pd.DataFrame(out)

    # ---------- Lineage ----------

    def build_lineage(self, tb_tree, recon_rows):
        by_key=defaultdict(list)
        for r in recon_rows:
            by_key[r.get('bs_mapping')].append(r)
        def attach(node, path):
            full = path if node['name']=='(Direct)' else path+[node['name']]
            key=' / '.join(full)
            node['matches']=by_key.get(key) or by_key.get(node['name']) or []
            for ch in node.get('children') or []:
                attach(ch, full)
            return node
        return [attach(g, []) for g in tb_tree]

    def flatten_lineage(self, tree, path=None):
        path=path or []; rows=[]
        for node in tree:
            cur = path if node['name']=='(Direct)' else path+[node['name']]
            matches=node.get('matches') or []
            children=node.get('children') or []
            if matches:
                for m in matches:
                    rows.append({'hierarchy':' / '.join(cur),'level':node['level'],'currency':m.get('currency'),
                                 'tb_amount':m.get('tb_amount'),'submission_amount':m.get('submission_amount'),
                                 'difference':m.get('difference'),'status':m.get('status'),'source':m.get('source'),
                                 'rule_type':m.get('rule_type'),'evidence':m.get('evidence')})
            elif not children:
                totals=node.get('currency_totals') or {}
                rows.append({'hierarchy':' / '.join(cur),'level':node['level'],'currency':'TOTAL',
                             'tb_amount':totals.get('TOTAL'),'submission_amount':None,'difference':None,
                             'status':'NO_RULE','source':None,'rule_type':None,'evidence':None})
            if children:
                rows.extend(self.flatten_lineage(children, cur).to_dict('records'))
        return pd.DataFrame(rows)

    # ---------- Export ----------

    def export(self, tb, pivot, submissions=None, recon=None, suggestions=None, lineage=None, path=None):
        path=path or tempfile.mktemp(suffix='.xlsx')
        with pd.ExcelWriter(path,engine='openpyxl') as w:
            summary=pd.DataFrame([{'Metric':'TB rows','Value':len(tb)},{'Metric':'TB adjusted balance','Value':tb.adjusted_balance.sum()},
                                   {'Metric':'BS Mapping groups','Value':tb.bs_mapping.nunique()},
                                   {'Metric':'Submission extracted lines','Value':0 if submissions is None else len(submissions)},
                                   {'Metric':'Reconciliation rules','Value':0 if recon is None else len(recon)}])
            summary.to_excel(w,sheet_name='01 Executive Summary',index=False); pivot.to_excel(w,sheet_name='02 TB Pivot',index=False); tb.to_excel(w,sheet_name='03 TB Account Detail',index=False)
            if submissions is not None and len(submissions): submissions.drop(columns=['hierarchy'],errors='ignore').to_excel(w,sheet_name='04 Submission Inventory',index=False)
            if recon is not None and len(recon): recon.to_excel(w,sheet_name='05 Recon Results',index=False)
            if suggestions: pd.DataFrame(suggestions).drop(columns=['components','match_basis'],errors='ignore').to_excel(w,sheet_name='06 Auto Suggestions',index=False)
            if lineage is not None and len(lineage): lineage.drop(columns=['evidence'],errors='ignore').to_excel(w,sheet_name='07 Reconciliation Map',index=False)
        wb=load_workbook(path)
        navy='173A5E'; blue='1D70B8'; light='EAF3FA'; thin=Side(style='thin',color='D9E2EA')
        for ws in wb.worksheets:
            ws.freeze_panes='A2'; ws.sheet_view.showGridLines=False; ws.auto_filter.ref=ws.dimensions
            for c in ws[1]: c.fill=PatternFill('solid',fgColor=navy); c.font=Font(color='FFFFFF',bold=True); c.alignment=Alignment(vertical='center')
            for col in ws.columns:
                letter=get_column_letter(col[0].column); width=min(42,max(12,max(len(str(x.value or '')) for x in col)+2)); ws.column_dimensions[letter].width=width
            for row in ws.iter_rows(min_row=2):
                for c in row:
                    if isinstance(c.value,(int,float)): c.number_format='#,##0.00;[Red](#,##0.00);-'
        wb.save(path); return path
